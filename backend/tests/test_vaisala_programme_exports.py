import csv
import io
import json

from openpyxl import load_workbook

from services.vaisala_programme import build_programme, DEFAULT_POLICY, ACTIONS
from services.vaisala_programme_exports import export_programme


def programme():
    result = build_programme([
        {"id": 1, "section_ref": "=HYPERLINK(\"https://invalid\")", "road_name": " École road",
         "assessment_scope": "section", "length_m": 100, "structural_pct": 2, "priority_score": 4},
        {"id": 2, "section_ref": "B", "road_name": "\t=BAD()", "assessment_scope": "section", "length_m": 100},
    ], survey_id=1, policy=DEFAULT_POLICY)
    result["generated_at"] = "2026-09-24T00:00:00+00:00"
    result["items"][0]["review"] = {"status": "deferred", "sequence": 1, "comment": " +SUM(1)", "assignee": "@TEST"}
    return result


def test_csv_unicode_formula_protection_full_payload_and_review():
    payload = programme()
    data = export_programme(payload, format="csv")
    rows = list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))
    assert len(rows) == 2
    structural = next(r for r in rows if r["section_ref"].startswith("'=HYPERLINK"))
    assert structural["section_ref"].startswith("'=")
    assert structural["road_name"] == " École road"
    # Locate review independently of engine display ordering.
    reviewed = next(r for r in rows if r["review_status"] == "deferred")
    assert reviewed["comment"] == "' +SUM(1)"
    assert reviewed["assignee"] == "'@TEST"
    assert json.loads(structural["source_evidence_json"])["structural_pct"] == 2
    assert all(r["policy_version"] == DEFAULT_POLICY["version"] for r in rows)


def test_xlsx_five_action_sheets_reconcile_and_preserve_empty_queues():
    payload = programme()
    payload["export_filters"] = {"filtered": True, "search": "B"}
    workbook = load_workbook(io.BytesIO(export_programme(payload, format="xlsx")))
    assert workbook.sheetnames == ["Summary", "All items", "Engineer assessment", "Evidence validation", "Treatment appraisal", "Monitor", "No action indicated", "Methodology", "Local defect review"]
    assert workbook["All items"].max_row - 1 == 2
    assert sum(workbook[name].max_row - 1 for name in workbook.sheetnames[2:7]) == 2
    assert workbook["Monitor"].max_row == 1
    assert all(cell.data_type != "f" for sheet in workbook for row in sheet for cell in row)
    notes = " ".join(str(row[0].value) for row in workbook["Methodology"])
    assert "Queue lengths may overlap" in notes
    summary = dict(workbook["Summary"].values)
    assert summary["Export scope"] == "filtered"
    assert '"search": "B"' in summary["Filters"]


def test_large_local_location_list_is_not_truncated_in_excel():
    payload = programme()
    item = payload['items'][0]
    flags = [dict(defect='Subsidence', location_status='Located', locations=[
        dict(net_reference='D1/1', from_m=i * 10, to_m=(i + 1) * 10) for i in range(1000)])]
    item['treatment_assessment']['local_defect_flags'] = flags
    workbook = load_workbook(io.BytesIO(export_programme(payload, format='xlsx')))
    assert workbook['Local defect review'].max_row == 1001
    parts = [row[3] for row in list(workbook['Evidence continuation'].values)[1:]
             if row[0] == item['item_key'] and row[1] == 'treatment_assessment']
    recovered = json.loads(''.join(parts))
    assert recovered['local_defect_flags'] == flags

import csv
import io
import json
import logging
from typing import Optional

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from database import get_db
from services.vaisala_qc import calculate_qc
from models.user import User
from models.vaisala import (
    VaisalaDefectWeight,
    VaisalaDefectWeightSet,
    VaisalaInterval,
    VaisalaNetworkFeature,
    VaisalaNetworkGeometry,
    VaisalaRagDriftLog,
    VaisalaRagThresholdSet,
    VaisalaSurvey,
    VaisalaSection,
)
from routers.auth import get_current_user, require_contributor, resolve_authority_id
from schemas.vaisala import (
    PaginatedVaisalaSections,
    VaisalaSectionOut,
    VaisalaStats,
    VaisalaSurveyOut,
    VaisalaUploadResult,
)
from services.vaisala_scoring import (
    NETWORK_CONFIGS,
    RAG_DERIVATION_NOTE,
    RAG_VALIDATED_WEIGHTS,
    assign_rag,
    assign_treatment,
    assign_treatment_percentile,
    detect_weight_drift,
    parse_raw_csv,
    parse_raw_xlsx,
    parse_shp_export,
)

router = APIRouter(prefix="/vaisala", tags=["vaisala"])
logger = logging.getLogger(__name__)


def _merge_feature_geometries(geometries: list):
    """Return the complete geometry represented by every feature for one section key."""
    from shapely.ops import unary_union

    return unary_union(geometries)


def _get_or_create_default_weight_set(db: Session) -> tuple:
    """Ensure the RAG-validated weight set and threshold set exist, return both IDs."""
    ws = db.query(VaisalaDefectWeightSet).filter_by(is_rag_validated=True).first()
    if not ws:
        ws = VaisalaDefectWeightSet(label="Default (RAG-validated)", is_rag_validated=True)
        db.add(ws)
        db.flush()
        for key, weight in RAG_VALIDATED_WEIGHTS.items():
            db.add(VaisalaDefectWeight(weight_set_id=ws.id, defect_key=key, weight=weight, active=True))
        ts = VaisalaRagThresholdSet(
            weight_set_id=ws.id,
            red_threshold=4.0,
            amber_threshold=1.8,
            derivation_note=RAG_DERIVATION_NOTE,
            is_active=True,
        )
        db.add(ts)
        db.flush()
    else:
        ts = db.query(VaisalaRagThresholdSet).filter_by(weight_set_id=ws.id, is_active=True).first()
        if not ts:
            ts = VaisalaRagThresholdSet(
                weight_set_id=ws.id,
                red_threshold=4.0,
                amber_threshold=1.8,
                derivation_note=RAG_DERIVATION_NOTE,
                is_active=True,
            )
            db.add(ts)
            db.flush()
    return ws.id, ts.id


def _persist_sections(db: Session, survey_id: int, sections: list[dict]) -> int:
    count = 0
    for s in sections:
        obj = VaisalaSection(survey_id=survey_id, **s)
        db.add(obj)
        count += 1
    return count


def _persist_intervals(db: Session, survey_id: int, intervals: list[dict]) -> int:
    BATCH = 2000
    count = 0
    for start in range(0, len(intervals), BATCH):
        batch = intervals[start:start + BATCH]
        db.bulk_insert_mappings(VaisalaInterval, [{"survey_id": survey_id, **iv} for iv in batch])
        count += len(batch)
    return count


def _interval_extras(iv) -> dict:
    if not iv.extras_json:
        return {}
    try:
        import json as _json
        return _json.loads(iv.extras_json)
    except Exception:
        return {}


def _interval_to_section_dict(iv, survey_id: int) -> dict:
    """Map a VaisalaInterval ORM row to a VaisalaSectionOut-shaped dict for the 10m raw view."""
    score = iv.interval_score or 0.0
    chunk_label = f"{iv.from_m:.0f}-{iv.to_m:.0f}m" if iv.from_m is not None and iv.to_m is not None else None
    return {
        "id": iv.id,
        "survey_id": survey_id,
        "section_ref": iv.section_ref,
        "chunk_label": chunk_label,
        "road_name": iv.road_name,
        "net_reference": iv.net_reference,
        "urban_rural": iv.urban_rural,
        "road_class": iv.road_class,
        "length_m": iv.length_m or 10.0,
        "priority_score": round(score, 4),
        "worst_interval_score": round(score, 4),
        "rag_band": assign_rag(score),
        "treatment": assign_treatment(iv.structural or 0.0, iv.alligator or 0.0, iv.localised or 0.0, iv.dressing or 0.0, iv.micro or 0.0, iv.edge or 0.0),
        "primary_defect": iv.primary_defect,
        "primary_defect_contribution": iv.primary_defect_contribution,
        "secondary_defect": None,
        "secondary_defect_contribution": None,
        "structural_pct": round((iv.structural or 0.0) * 100, 2),
        "localised_pct":  round((iv.localised  or 0.0) * 100, 2),
        "dressing_pct":   round((iv.dressing   or 0.0) * 100, 2),
        "micro_pct":      round((iv.micro      or 0.0) * 100, 2),
        "alligator_pct":  round((iv.alligator  or 0.0) * 100, 2),
        "edge_pct":       round((iv.edge       or 0.0) * 100, 2),
        "road_surface_condition": iv.road_surface_condition,
        "road_surface_condition_class": iv.road_surface_condition_class,
        "asphalt_condition": iv.asphalt_condition,
        "asphalt_condition_class": iv.asphalt_condition_class,
        "pas2161_category": iv.pas2161_category,
        **calculate_qc([(iv.length_m or 10.0, _interval_extras(iv))]),
        "extras": _interval_extras(iv),
    }


def _chunk_by_length(ivs_sorted: list, target_len: float) -> list[list]:
    """Emit chunks of ~target_len meters. Trailing remainder < half target merged into last chunk."""
    chunks: list[list] = []
    current: list = []
    current_len = 0.0
    for iv in ivs_sorted:
        current.append(iv)
        current_len += iv.length_m or 10.0
        if current_len >= target_len:
            chunks.append(current)
            current = []
            current_len = 0.0
    if current:
        if chunks and current_len < (target_len / 2):
            chunks[-1].extend(current)
        else:
            chunks.append(current)
    return chunks


def _chunk_respecting_subgroups(ivs: list, target_len: float) -> list[list]:
    """When a section spans multiple net_reference sub-stretches, chunk each independently.

    Prevents 100m windows from mixing intervals across physically disconnected stretches
    (documented WSCC case: NSGNO 45100793 fan-shaped geometry across 7 WSCCNET sub-stretches).
    """
    subs: dict = {}
    for iv in ivs:
        k = iv.net_reference or "__none__"
        subs.setdefault(k, []).append(iv)
    if len(subs) <= 1:
        ivs_sorted = sorted(ivs, key=lambda x: (x.from_m is None, x.from_m or 0.0))
        return _chunk_by_length(ivs_sorted, target_len)
    all_chunks: list[list] = []
    for group in subs.values():
        sorted_group = sorted(group, key=lambda x: (x.from_m is None, x.from_m or 0.0))
        all_chunks.extend(_chunk_by_length(sorted_group, target_len))
    all_chunks.sort(key=lambda c: (c[0].from_m is None, c[0].from_m or 0.0))
    return all_chunks


def _intervals_to_100m_sections(intervals: list, survey_id: int) -> list[dict]:
    """Group intervals into ~100m windows per section and return section-level dicts."""
    from collections import defaultdict

    by_section: dict = defaultdict(list)
    for iv in intervals:
        by_section[iv.section_ref].append(iv)

    result = []
    chunk_counter = 0

    for section_ref, ivs in by_section.items():
        chunks = _chunk_respecting_subgroups(ivs, 100.0)
        n_chunks = len(chunks)

        for i, chunk in enumerate(chunks):
            chunk_counter += 1
            total_len = sum(iv.length_m or 10.0 for iv in chunk)
            if total_len <= 0:
                continue
            w = [(iv.length_m or 10.0) / total_len for iv in chunk]

            score   = sum(w[i] * (chunk[i].interval_score or 0.0) for i in range(len(chunk)))
            worst   = max(iv.interval_score or 0.0 for iv in chunk)
            struct  = sum(w[i] * (chunk[i].structural  or 0.0) for i in range(len(chunk)))
            allig   = sum(w[i] * (chunk[i].alligator   or 0.0) for i in range(len(chunk)))
            local   = sum(w[i] * (chunk[i].localised   or 0.0) for i in range(len(chunk)))
            dress   = sum(w[i] * (chunk[i].dressing    or 0.0) for i in range(len(chunk)))
            micro_v = sum(w[i] * (chunk[i].micro       or 0.0) for i in range(len(chunk)))
            edge    = sum(w[i] * (chunk[i].edge        or 0.0) for i in range(len(chunk)))

            rsc_pairs  = [(chunk[i].road_surface_condition, w[i]) for i in range(len(chunk)) if chunk[i].road_surface_condition is not None]
            rsc        = sum(v * ww for v, ww in rsc_pairs) / sum(ww for _, ww in rsc_pairs) if rsc_pairs else None
            asph_pairs = [(chunk[i].asphalt_condition, w[i]) for i in range(len(chunk)) if chunk[i].asphalt_condition is not None]
            asph       = sum(v * ww for v, ww in asph_pairs) / sum(ww for _, ww in asph_pairs) if asph_pairs else None
            pas_vals   = [iv.pas2161_category for iv in chunk if iv.pas2161_category]
            pas        = max(pas_vals) if pas_vals else None

            best   = max(chunk, key=lambda x: x.interval_score or 0.0)
            from_p = chunk[0].from_m
            to_p   = chunk[-1].to_m
            chunk_label = f"{from_p:.0f}-{to_p:.0f}m (chunk {i+1}/{n_chunks})" if from_p is not None and to_p is not None else f"chunk {i+1}/{n_chunks}"

            result.append({
                "id": -chunk_counter,
                "survey_id": survey_id,
                "section_ref": section_ref,
                "chunk_label": chunk_label,
                "road_name": chunk[0].road_name,
                "net_reference": chunk[0].net_reference,
                "urban_rural": chunk[0].urban_rural,
                "road_class": chunk[0].road_class,
                "length_m": total_len,
                "priority_score": round(score, 4),
                "worst_interval_score": round(worst, 4),
                "rag_band": assign_rag(score),
                "treatment": assign_treatment(struct, allig, local, dress, micro_v, edge),
                "primary_defect": best.primary_defect,
                "primary_defect_contribution": best.primary_defect_contribution,
                "secondary_defect": None,
                "secondary_defect_contribution": None,
                "structural_pct": round(struct * 100, 2),
                "localised_pct":  round(local  * 100, 2),
                "dressing_pct":   round(dress  * 100, 2),
                "micro_pct":      round(micro_v * 100, 2),
                "alligator_pct":  round(allig  * 100, 2),
                "edge_pct":       round(edge   * 100, 2),
                "road_surface_condition": rsc,
                "road_surface_condition_class": None,
                "asphalt_condition": asph,
                "asphalt_condition_class": None,
                "pas2161_category": pas,
                **calculate_qc((iv.length_m or 10.0, _interval_extras(iv)) for iv in chunk),
                # priority_dst.html keeps "Video Link (start of extent)" at 100m; capture first
                # interval's link-ish extras only (aggregate view can't sensibly represent every
                # sub-interval's link, so start-of-extent is the ref convention)
                "extras": _interval_extras(chunk[0]) if chunk else {},
            })

    return result


_SECTION_OUT_FIELDS = [
    "id", "survey_id", "section_ref", "road_name", "net_reference", "urban_rural", "road_class",
    "length_m", "priority_score", "worst_interval_score", "rag_band", "treatment",
    "primary_defect", "primary_defect_contribution",
    "secondary_defect", "secondary_defect_contribution",
    "structural_pct", "localised_pct", "dressing_pct", "micro_pct",
    "alligator_pct", "edge_pct",
    "road_surface_condition", "road_surface_condition_class",
    "asphalt_condition", "asphalt_condition_class", "pas2161_category",
    "qc_completeness_pct", "qc_completeness_band",
    "qc_reliability_pct", "qc_reliability_band",
    "chunk_label",
]


def _to_dict_row(row) -> dict:
    if isinstance(row, dict):
        if "extras" not in row:
            row = {**row, "extras": {}}
        return row
    d = {f: getattr(row, f, None) for f in _SECTION_OUT_FIELDS}
    d["extras"] = {}
    return d


def _validate_view_params(merge_scale: str, split: str, treatment_mode: str) -> None:
    if merge_scale not in ("section", "100m", "10m"):
        raise HTTPException(400, "merge_scale must be one of: section, 100m, 10m")
    if split not in ("combined", "urban", "rural"):
        raise HTTPException(400, "split must be one of: combined, urban, rural")
    if treatment_mode not in ("defect", "percentile"):
        raise HTTPException(400, "treatment_mode must be one of: defect, percentile")


def _db_sections_as_dicts(db: Session, survey_id: int, urban_rural: Optional[str] = None) -> list[dict]:
    q = db.query(VaisalaSection).filter_by(survey_id=survey_id)
    if urban_rural is not None:
        q = q.filter(VaisalaSection.urban_rural == urban_rural)
    return [_to_dict_row(s) for s in q.all()]


def _build_view_rows(db: Session, survey_id: int, merge_scale: str, split: str) -> list[dict]:
    """Return dict rows for the requested view.

    Urban rows always use section scale (per priority_dst.html brief: 'urban roads always score
    at whole section length regardless of the selected scale'). Combined view at 10m/100m
    concatenates urban-at-section + rural-at-selected-scale.
    """
    if split == "urban":
        return _db_sections_as_dicts(db, survey_id, urban_rural="U")

    if split == "rural":
        if merge_scale == "section":
            return _db_sections_as_dicts(db, survey_id, urban_rural="R")
        scaled = _scaled_section_rows(db, survey_id, merge_scale)
        if not scaled:
            raise HTTPException(422, "No interval data stored for this survey. Re-upload the raw file to enable 10m/100m scales.")
        return [r for r in scaled if r.get("urban_rural") == "R"]

    # combined
    if merge_scale == "section":
        return _db_sections_as_dicts(db, survey_id)
    scaled = _scaled_section_rows(db, survey_id, merge_scale)
    if not scaled:
        raise HTTPException(422, "No interval data stored for this survey. Re-upload the raw file to enable 10m/100m scales.")
    urban_section = _db_sections_as_dicts(db, survey_id, urban_rural="U")
    rural_scaled = [r for r in scaled if r.get("urban_rural") == "R"]
    if not urban_section:
        # network has no U/R indicator — just return scaled rows as-is
        return scaled
    return urban_section + rural_scaled


def _apply_percentile_treatments(rows: list[dict]) -> None:
    """Overwrite each row's treatment based on percentile rank of priority_score across the full set.

    Mirrors priority_dst.html assignTreatmentsByPercentile: sort ascending by score, rank 100 = worst,
    then choose band from PERCENTILE_TREATMENTS. Rows without a numeric score get '—'.
    """
    scored = [r for r in rows if r.get("priority_score") is not None]
    unscored = [r for r in rows if r.get("priority_score") is None]
    scored.sort(key=lambda r: r["priority_score"])
    n = len(scored)
    for idx, r in enumerate(scored):
        pct = (idx / (n - 1)) * 100 if n > 1 else 100
        r["treatment"] = assign_treatment_percentile(pct)
    for r in unscored:
        r["treatment"] = "—"


def _scaled_section_rows(db: Session, survey_id: int, merge_scale: str) -> list[dict]:
    """Return VaisalaSectionOut-shaped dicts at the requested scale. Empty list if no intervals stored."""
    intervals = db.query(VaisalaInterval).filter_by(survey_id=survey_id).all()
    if not intervals:
        return []
    if merge_scale == "100m":
        return _intervals_to_100m_sections(intervals, survey_id)
    return [_interval_to_section_dict(iv, survey_id) for iv in intervals]


def _filter_and_sort_rows(
    rows: list[dict],
    rag_band: Optional[str],
    treatment: Optional[str],
    sort_by: str,
    sort_dir: str,
) -> list[dict]:
    if rag_band:
        rows = [r for r in rows if r.get("rag_band") == rag_band]
    if treatment:
        rows = [r for r in rows if r.get("treatment") == treatment]
    valid = [r for r in rows if r.get(sort_by) is not None]
    nulls = [r for r in rows if r.get(sort_by) is None]
    valid.sort(key=lambda r: r.get(sort_by), reverse=(sort_dir != "asc"))
    return valid + nulls


def _rag_summary(sections: list[dict]) -> dict:
    summary = {"Red": 0, "Amber": 0, "Green": 0}
    for s in sections:
        band = s.get("rag_band") or "Green"
        summary[band] = summary.get(band, 0) + 1
    return summary


@router.post("/upload/raw", response_model=VaisalaUploadResult, status_code=201)
async def upload_raw(
    file: UploadFile = File(...),
    network_key: str = Query("stroud", description="stroud or wscc"),
    dedup_strategy: str = Query("latest", description="latest | none — resolve duplicate survey passes per (section,from,to)"),
    weights_json: Optional[str] = Query(None, description="JSON object of defect weights; omit to use RAG-validated defaults"),
    current_user: User = Depends(require_contributor),
    db: Session = Depends(get_db),
):
    """Upload a raw Vaisala interval-level XLSX or CSV file. Scores server-side."""
    if network_key not in NETWORK_CONFIGS:
        raise HTTPException(400, f"Unknown network_key '{network_key}'. Valid: {list(NETWORK_CONFIGS)}")
    if dedup_strategy not in ("latest", "none"):
        raise HTTPException(400, "dedup_strategy must be 'latest' or 'none'")

    custom_weights: Optional[dict] = None
    if weights_json:
        try:
            custom_weights = json.loads(weights_json)
            if not isinstance(custom_weights, dict):
                raise ValueError("weights must be a JSON object")
            custom_weights = {k: float(v) for k, v in custom_weights.items()}
        except (ValueError, TypeError) as e:
            raise HTTPException(400, f"Invalid weights_json: {e}")

    content = await file.read()
    filename = file.filename or "upload"
    ext = filename.rsplit(".", 1)[-1].lower()

    try:
        if ext in ("xlsx", "xls"):
            result = parse_raw_xlsx(content, network_key, weights=custom_weights, dedup_strategy=dedup_strategy)
            fmt = "xlsx"
        elif ext == "csv":
            result = parse_raw_csv(content, network_key, weights=custom_weights, dedup_strategy=dedup_strategy)
            fmt = "csv"
        else:
            raise HTTPException(400, "Unsupported file type. Upload XLSX, XLS, or CSV.")
    except ValueError as e:
        raise HTTPException(422, str(e))
    except Exception as e:
        logger.exception("Vaisala raw parse failed")
        raise HTTPException(500, f"Parse error: {e}")

    sections = result["sections"]
    if not sections:
        raise HTTPException(422, "No sections could be scored from this file. Check the network selection and file format.")

    intervals = result.get("intervals", [])
    drift = result["drift"]
    dup_groups = result.get("dup_groups", 0)
    info_meta = result.get("info_meta", {}) or {}
    flattened_link_warnings = result.get("flattened_link_warnings", []) or []
    ws_id, ts_id = _get_or_create_default_weight_set(db)

    notes_payload = json.dumps({"info_meta": info_meta, "dup_groups_resolved": dup_groups, "dedup_strategy": dedup_strategy})

    survey = VaisalaSurvey(
        authority_id=current_user.authority_id,
        source_filename=filename,
        source_format=fmt,
        network_key=network_key,
        weight_set_id=ws_id,
        threshold_set_id=ts_id,
        row_count=result["row_count"],
        section_count=len(sections),
        has_weight_drift=len(drift) > 0,
        notes=notes_payload,
    )
    db.add(survey)
    db.flush()

    for d in drift:
        db.add(VaisalaRagDriftLog(
            survey_id=survey.id,
            defect_key=d["defect_key"],
            validated_weight=d["validated_weight"],
            actual_weight=d["actual_weight"],
        ))

    _persist_sections(db, survey.id, sections)
    if intervals:
        _persist_intervals(db, survey.id, intervals)
    db.commit()

    return VaisalaUploadResult(
        survey_id=survey.id,
        source_filename=filename,
        network_key=network_key,
        row_count=result["row_count"],
        section_count=len(sections),
        has_weight_drift=len(drift) > 0,
        drift_details=drift,
        rag_summary=_rag_summary(sections),
        dup_groups_resolved=dup_groups,
        info_meta=info_meta,
        flattened_link_warnings=flattened_link_warnings,
    )


@router.post("/upload/shp", response_model=VaisalaUploadResult, status_code=201)
async def upload_shp_export(
    file: UploadFile = File(...),
    network_key: str = Query("stroud"),
    current_user: User = Depends(require_contributor),
    db: Session = Depends(get_db),
):
    """Upload a zipped SHP export from priority_dst.html (pre-scored sections)."""
    content = await file.read()
    filename = file.filename or "export.zip"

    try:
        result = parse_shp_export(content)
    except Exception as e:
        logger.exception("Vaisala SHP parse failed")
        raise HTTPException(422, f"Could not read SHP export: {e}")

    sections = result["sections"]
    if not sections:
        raise HTTPException(422, "No sections found in SHP export.")

    ws_id, ts_id = _get_or_create_default_weight_set(db)

    survey = VaisalaSurvey(
        authority_id=current_user.authority_id,
        source_filename=filename,
        source_format="shp",
        network_key=network_key,
        weight_set_id=ws_id,
        threshold_set_id=ts_id,
        row_count=len(sections),
        section_count=len(sections),
        has_weight_drift=False,
    )
    db.add(survey)
    db.flush()

    _persist_sections(db, survey.id, sections)
    db.commit()

    return VaisalaUploadResult(
        survey_id=survey.id,
        source_filename=filename,
        network_key=network_key,
        row_count=len(sections),
        section_count=len(sections),
        has_weight_drift=False,
        drift_details=[],
        rag_summary=_rag_summary(sections),
    )


def _get_survey(db: Session, survey_id: int, current_user) -> Optional[VaisalaSurvey]:
    """Load survey by ID; non-admins must own the survey."""
    q = db.query(VaisalaSurvey).filter(VaisalaSurvey.id == survey_id)
    if current_user.role != "admin":
        q = q.filter(VaisalaSurvey.authority_id == current_user.authority_id)
    return q.first()


@router.get("/surveys", response_model=list[VaisalaSurveyOut])
def list_surveys(
    authority_id: Optional[int] = Query(None, description="Admin only — filter by authority"),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    auth_id = resolve_authority_id(current_user, authority_id)
    q = db.query(VaisalaSurvey)
    if auth_id is not None:
        q = q.filter(VaisalaSurvey.authority_id == auth_id)
    return q.order_by(VaisalaSurvey.imported_at.desc()).all()


@router.get("/surveys/{survey_id}/stats", response_model=VaisalaStats)
def survey_stats(
    survey_id: int,
    merge_scale: str = Query("section", description="section | 100m | 10m"),
    split: str = Query("combined", description="combined | urban | rural"),
    treatment_mode: str = Query("defect", description="defect | percentile"),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    survey = _get_survey(db, survey_id, current_user)
    if not survey:
        raise HTTPException(404, "Survey not found")

    _validate_view_params(merge_scale, split, treatment_mode)
    effective_scale = "section" if split == "urban" else merge_scale

    rows = _build_view_rows(db, survey_id, effective_scale, split)
    if treatment_mode == "percentile":
        _apply_percentile_treatments(rows)

    total_length_km = sum((r.get("length_m") or 0) for r in rows) / 1000
    rag_counts: dict = {"Red": 0, "Amber": 0, "Green": 0}
    rag_length: dict = {"Red": 0.0, "Amber": 0.0, "Green": 0.0}
    treatment_counts: dict = {}

    for r in rows:
        b = r.get("rag_band") or "Green"
        rag_counts[b] = rag_counts.get(b, 0) + 1
        rag_length[b] = rag_length.get(b, 0.0) + (r.get("length_m") or 0) / 1000
        t = r.get("treatment") or "Monitor / Patching"
        treatment_counts[t] = treatment_counts.get(t, 0) + 1

    has_urban_rural = (
        db.query(VaisalaSection.id)
          .filter(VaisalaSection.survey_id == survey_id, VaisalaSection.urban_rural.isnot(None))
          .first()
        is not None
    )

    return VaisalaStats(
        survey_id=survey.id,
        source_filename=survey.source_filename,
        network_key=survey.network_key,
        imported_at=survey.imported_at,
        section_count=len(rows),
        total_length_km=round(total_length_km, 2),
        rag_counts=rag_counts,
        rag_length_km={k: round(v, 2) for k, v in rag_length.items()},
        has_weight_drift=survey.has_weight_drift,
        top_treatments=treatment_counts,
        has_urban_rural=has_urban_rural,
    )


@router.get("/surveys/{survey_id}/sections", response_model=PaginatedVaisalaSections)
def list_sections(
    survey_id: int,
    rag_band: Optional[str] = Query(None),
    treatment: Optional[str] = Query(None),
    skip: int = Query(0, ge=0),
    limit: int = Query(50, ge=1, le=200),
    sort_by: str = Query("priority_score"),
    sort_dir: str = Query("desc"),
    merge_scale: str = Query("section", description="section | 100m | 10m"),
    split: str = Query("combined", description="combined | urban | rural"),
    treatment_mode: str = Query("defect", description="defect | percentile"),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    survey = _get_survey(db, survey_id, current_user)
    if not survey:
        raise HTTPException(404, "Survey not found")

    _validate_view_params(merge_scale, split, treatment_mode)
    effective_scale = "section" if split == "urban" else merge_scale

    rows = _build_view_rows(db, survey_id, effective_scale, split)
    if treatment_mode == "percentile":
        _apply_percentile_treatments(rows)

    rows = _filter_and_sort_rows(rows, rag_band, treatment, sort_by, sort_dir)
    total = len(rows)
    page = rows[skip:skip + limit]
    return PaginatedVaisalaSections(sections=page, total=total, survey_id=survey_id)


@router.get("/surveys/{survey_id}/sections/all", response_model=list[VaisalaSectionOut])
def all_sections(
    survey_id: int,
    merge_scale: str = Query("section", description="section | 100m | 10m"),
    split: str = Query("combined", description="combined | urban | rural"),
    treatment_mode: str = Query("defect", description="defect | percentile"),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Return all rows for a survey without pagination — used for correlation and QC analysis."""
    survey = _get_survey(db, survey_id, current_user)
    if not survey:
        raise HTTPException(404, "Survey not found")

    _validate_view_params(merge_scale, split, treatment_mode)
    effective_scale = "section" if split == "urban" else merge_scale
    rows = _build_view_rows(db, survey_id, effective_scale, split)
    if treatment_mode == "percentile":
        _apply_percentile_treatments(rows)
    return rows


_LINK_HDR_PATTERN = __import__("re").compile(r"link|url|hyperlink|video", __import__("re").IGNORECASE)
RAG_FILL = {"Red": "FFC0432F", "Amber": "FFD9A51C", "Green": "FF3A7D44"}
TREATMENT_FILL = {
    "Resurfacing":       "FFC0432F",
    "Patching":          "FFD9862A",
    "Surface Dressing":  "FFD9A51C",
    "Micro-surfacing":   "FF7A9B6E",
    "Monitor / Patching":"FF4C6B6F",
}


@router.get("/surveys/{survey_id}/export")
def export_csv(
    survey_id: int,
    merge_scale: str = Query("section", description="section | 100m | 10m"),
    split: str = Query("combined", description="combined | urban | rural"),
    treatment_mode: str = Query("defect", description="defect | percentile"),
    format: str = Query("csv", description="csv | xlsx | shp"),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    survey = _get_survey(db, survey_id, current_user)
    if not survey:
        raise HTTPException(404, "Survey not found")

    _validate_view_params(merge_scale, split, treatment_mode)
    if format not in ("csv", "xlsx", "shp"):
        raise HTTPException(400, "format must be one of: csv, xlsx, shp")
    effective_scale = "section" if split == "urban" else merge_scale

    # Human-readable column layout (label -> row key), mirroring priority_dst.html exportData().
    # Defect-proportion columns only meaningful when treatment_mode='defect' (per ref: those
    # proportions are what the defect-pattern algorithm consumed; percentile mode ignores them).
    base_layout: list[tuple[str, str]] = [
        ("Section",                        "section_ref"),
        ("Extent",                         "chunk_label"),
        ("Road Name",                      "road_name"),
        ("Net Reference",                  "net_reference"),
        ("Urban/Rural",                    "urban_rural"),
        ("Road Class",                     "road_class"),
        ("Length (m)",                     "length_m"),
        ("Road Surface Condition",         "road_surface_condition"),
        ("Road Surface Condition Class",   "road_surface_condition_class"),
        ("Asphalt Condition",              "asphalt_condition"),
        ("Asphalt Condition Class",        "asphalt_condition_class"),
        ("PAS 2161 RCM Category",          "pas2161_category"),
        ("List 4 Weighted Score",          "priority_score"),
        ("List 4 Worst Interval Score",    "worst_interval_score"),
        ("RAG Band",                       "rag_band"),
        ("List 4 Treatment",               "treatment"),
        ("Primary Defect",                 "primary_defect"),
        ("Primary Defect Contribution",    "primary_defect_contribution"),
        ("Secondary Defect",               "secondary_defect"),
        ("Secondary Defect Contribution",  "secondary_defect_contribution"),
    ]
    defect_layout: list[tuple[str, str]] = [
        ("Structural Defect Proportion (%)",         "structural_pct"),
        ("Alligator Cracking Proportion (%)",        "alligator_pct"),
        ("Localised Defect Proportion (%)",          "localised_pct"),
        ("Edge Deterioration Proportion (%)",        "edge_pct"),
        ("Surface (Dressing-type) Proportion (%)",   "dressing_pct"),
        ("Surface (Micro-type) Proportion (%)",      "micro_pct"),
    ]
    qc_layout: list[tuple[str, str]] = [
        ("Survey Completeness (%)",       "qc_completeness_pct"),
        ("Survey Completeness Band",      "qc_completeness_band"),
        ("Reading Reliability (%)",       "qc_reliability_pct"),
        ("Reading Reliability Band",      "qc_reliability_band"),
    ]
    layout = base_layout + (defect_layout if treatment_mode == "defect" else []) + qc_layout

    rows = _build_view_rows(db, survey_id, effective_scale, split)
    if treatment_mode == "percentile":
        _apply_percentile_treatments(rows)
    rows.sort(key=lambda r: (r.get("priority_score") is None, -(r.get("priority_score") or 0.0)))

    # Raw pass-through columns from recovered hyperlinks + coord/time fields — only meaningful
    # at 10m raw scale, or first-of-extent at 100m (see _intervals_to_100m_sections). Section
    # scale skips these entirely (DB sections have no per-row raw source columns).
    extras_labels: list[str] = []
    if effective_scale in ("10m", "100m"):
        seen: dict[str, None] = {}
        for r in rows:
            for k in (r.get("extras") or {}).keys():
                if k not in seen:
                    seen[k] = None
                    extras_labels.append(k)

    labels = [label for label, _ in layout] + extras_labels

    safe_name = survey.source_filename.rsplit(".", 1)[0].replace(" ", "_")
    scale_part = "" if merge_scale == "section" else f"_{merge_scale}"
    split_part = "" if split == "combined" else f"_{split}"
    tmode_part = "" if treatment_mode == "defect" else f"_{treatment_mode}"
    filename_base = f"vaisala_{safe_name}{scale_part}{split_part}{tmode_part}"

    def build_row_out(s: dict) -> dict:
        row_out = {label: s.get(key) for label, key in layout}
        extras = s.get("extras") or {}
        for k in extras_labels:
            row_out[k] = extras.get(k)
        return row_out

    if format == "csv":
        def generate_csv():
            buf = io.StringIO()
            w = csv.DictWriter(buf, fieldnames=labels)
            w.writeheader()
            yield buf.getvalue()
            for s in rows:
                buf = io.StringIO()
                w = csv.DictWriter(buf, fieldnames=labels)
                w.writerow(build_row_out(s))
                yield buf.getvalue()

        return StreamingResponse(
            generate_csv(),
            media_type="text/csv",
            headers={"Content-Disposition": f'attachment; filename="{filename_base}.csv"'},
        )

    # xlsx path
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = "Priority List"
    ws.append(labels)
    header = ws[1]
    header_fill = PatternFill(fill_type="solid", start_color="FF262A28", end_color="FF262A28")
    header_font = Font(bold=True, color="FFE9E6DC")
    for cell in header:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="left", vertical="center")

    rag_col_idx = labels.index("RAG Band") + 1 if "RAG Band" in labels else None
    tmt_col_idx = labels.index("List 4 Treatment") + 1 if "List 4 Treatment" in labels else None
    link_col_idxs = [labels.index(l) + 1 for l in extras_labels if _LINK_HDR_PATTERN.search(l)]

    for s in rows:
        row_dict = build_row_out(s)
        ws.append([row_dict.get(l) for l in labels])
        row_num = ws.max_row
        if rag_col_idx:
            band = row_dict.get("RAG Band")
            argb = RAG_FILL.get(band)
            if argb:
                ws.cell(row=row_num, column=rag_col_idx).fill = PatternFill(fill_type="solid", start_color=argb, end_color=argb)
                ws.cell(row=row_num, column=rag_col_idx).font = Font(bold=True, color="FFFFFFFF")
        if tmt_col_idx:
            t = row_dict.get("List 4 Treatment")
            argb = TREATMENT_FILL.get(t)
            if argb:
                ws.cell(row=row_num, column=tmt_col_idx).fill = PatternFill(fill_type="solid", start_color=argb, end_color=argb)
                ws.cell(row=row_num, column=tmt_col_idx).font = Font(color="FFFFFFFF")
        for col_idx in link_col_idxs:
            cell = ws.cell(row=row_num, column=col_idx)
            v = cell.value
            if isinstance(v, str) and v.lower().startswith("http"):
                cell.hyperlink = v
                cell.font = Font(color="FF3B78E7", underline="single")

    ws.freeze_panes = "A2"
    for col_letter in [c[0].column_letter for c in ws.iter_cols(min_row=1, max_row=1)]:
        ws.column_dimensions[col_letter].width = 22

    if format == "xlsx":
        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)
        return StreamingResponse(
            iter([stream.getvalue()]),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename_base}.xlsx"'},
        )

    # ─── SHP export ───
    # Requires an active network geometry layer for real road geometry. Row scores are joined
    # by section_ref → section_key. Rows without a matching feature are dropped (row count in
    # response header helps the frontend surface a "missed N rows" note).
    geom = (
        db.query(VaisalaNetworkGeometry)
          .filter_by(authority_id=current_user.authority_id, is_active=True)
          .order_by(VaisalaNetworkGeometry.created_at.desc())
          .first()
    )
    if not geom:
        raise HTTPException(422, "SHP export needs an active network geometry — upload a network shapefile first.")

    def _norm(k: str) -> str:
        s = str(k)
        return (s.lstrip("0") or "0") if s.isdigit() else s

    feature_rows = db.query(VaisalaNetworkFeature).filter_by(geometry_id=geom.id).all()
    import json as _json
    from shapely.geometry import shape as _shape
    features_by_key: dict = {}
    features_by_norm: dict = {}
    for f in feature_rows:
        try:
            shp = _shape(_json.loads(f.geometry_geojson))
        except Exception:
            continue
        features_by_key.setdefault(f.section_key, []).append(shp)
        features_by_norm.setdefault(_norm(f.section_key), []).append(shp)

    # SHP DBF column names capped at 10 chars — use short forms.
    short_layout = [
        ("Section",    "section_ref"),
        ("Extent",     "chunk_label"),
        ("Road",       "road_name"),
        ("NetRef",     "net_reference"),
        ("UrbRur",     "urban_rural"),
        ("RoadClass",  "road_class"),
        ("Length_m",   "length_m"),
        ("Score",      "priority_score"),
        ("WorstScore", "worst_interval_score"),
        ("RAG",        "rag_band"),
        ("Treatment",  "treatment"),
        ("PrimDefect", "primary_defect"),
        ("PrimContr",  "primary_defect_contribution"),
        ("StructPct",  "structural_pct"),
        ("AlligPct",   "alligator_pct"),
        ("LocalPct",   "localised_pct"),
        ("DressPct",   "dressing_pct"),
        ("MicroPct",   "micro_pct"),
        ("EdgePct",    "edge_pct"),
    ]

    import geopandas as _gpd
    records = []
    geoms = []
    skipped = 0
    for r in rows:
        ref = str(r.get("section_ref") or "")
        matched_geometries = features_by_key.get(ref)
        if matched_geometries is None:
            matched_geometries = features_by_norm.get(_norm(ref))
        if not matched_geometries:
            skipped += 1
            continue
        rec = {short: r.get(key) for short, key in short_layout}
        records.append(rec)
        geoms.append(_merge_feature_geometries(matched_geometries))

    if not records:
        raise HTTPException(422, f"No rows matched the network geometry (checked {len(rows)} rows against {len(feature_rows)} features). Confirm the section-field on upload matches how sections are labelled.")

    gdf = _gpd.GeoDataFrame(records, geometry=geoms, crs="EPSG:4326")

    import tempfile as _tempfile
    import zipfile as _zipfile
    import os as _os
    with _tempfile.TemporaryDirectory() as tmp:
        shp_out = _os.path.join(tmp, "sections.shp")
        try:
            gdf.to_file(shp_out, driver="ESRI Shapefile")
        except Exception as e:
            logger.exception("SHP write failed")
            raise HTTPException(500, f"Could not write shapefile: {e}")

        stream = io.BytesIO()
        with _zipfile.ZipFile(stream, "w", _zipfile.ZIP_DEFLATED) as zf:
            for f in _os.listdir(tmp):
                zf.write(_os.path.join(tmp, f), arcname=f)
        stream.seek(0)

    return StreamingResponse(
        iter([stream.getvalue()]),
        media_type="application/zip",
        headers={
            "Content-Disposition": f'attachment; filename="{filename_base}.zip"',
            "X-Rows-Skipped-Unmatched": str(skipped),
            "X-Rows-Written": str(len(records)),
        },
    )


# ============================================================================
# Network geometry (client's own road network SHP) — used by Map tab basemap join.
# Mirrors priority_dst.html readNetworkShp / dbf logic but done server-side via
# geopandas + pyogrio. Features stored per-section-key as WGS84 GeoJSON so the
# frontend Leaflet layer can render them directly with OS Maps as the basemap.
# ============================================================================

@router.post("/network-geometry/upload", status_code=201)
async def upload_network_geometry(
    file: UploadFile = File(..., description="ZIP containing .shp/.shx/.dbf/.prj"),
    section_field: Optional[str] = Query(None, description="DBF column identifying each section (auto-detects if omitted)"),
    current_user: User = Depends(require_contributor),
    db: Session = Depends(get_db),
):
    import io as _io
    import os as _os
    import tempfile as _tempfile
    import zipfile as _zipfile

    content = await file.read()
    filename = file.filename or "network.zip"
    if not filename.lower().endswith(".zip"):
        raise HTTPException(400, "Upload a ZIP containing the .shp/.shx/.dbf/.prj set")

    with _tempfile.TemporaryDirectory() as tmp:
        try:
            with _zipfile.ZipFile(_io.BytesIO(content)) as zf:
                zf.extractall(tmp)
        except _zipfile.BadZipFile:
            raise HTTPException(422, "Uploaded file is not a valid ZIP")

        shp_files = []
        for root, _dirs, files in _os.walk(tmp):
            for f in files:
                if f.lower().endswith(".shp"):
                    shp_files.append(_os.path.join(root, f))
        if not shp_files:
            raise HTTPException(422, "ZIP does not contain a .shp file")
        shp_path = shp_files[0]

        try:
            import geopandas as gpd
            gdf = gpd.read_file(shp_path)
        except Exception as e:
            logger.exception("Network SHP read failed")
            raise HTTPException(422, f"Could not read shapefile: {e}")

    if gdf.empty:
        raise HTTPException(422, "Shapefile contained zero features")

    fields = [c for c in gdf.columns if c != "geometry"]
    if section_field is None:
        # Auto-detect: prefer SECTIONLAB, then anything containing 'section' or 'nsg'
        upper_fields = {c.upper(): c for c in fields}
        candidate = (
            upper_fields.get("SECTIONLAB")
            or next((c for c in fields if "section" in c.lower()), None)
            or next((c for c in fields if "nsg" in c.lower()), None)
            or (fields[0] if fields else None)
        )
        if not candidate:
            raise HTTPException(422, "No usable attribute field for section key")
        section_field = candidate
    if section_field not in fields:
        raise HTTPException(400, f"section_field '{section_field}' not present in DBF (available: {fields[:12]})")

    crs_wkt = None
    try:
        if gdf.crs is not None:
            crs_wkt = gdf.crs.to_wkt()
            if str(gdf.crs) != "EPSG:4326":
                gdf = gdf.to_crs("EPSG:4326")
    except Exception:
        pass

    # Deactivate any previous active layer for this authority (only one active at a time)
    db.query(VaisalaNetworkGeometry).filter_by(authority_id=current_user.authority_id, is_active=True).update({"is_active": False})

    geom_row = VaisalaNetworkGeometry(
        authority_id=current_user.authority_id,
        source_filename=filename,
        section_field=section_field,
        feature_count=int(len(gdf)),
        crs_wkt=crs_wkt,
        is_active=True,
    )
    db.add(geom_row)
    db.flush()

    import json as _json
    from shapely.geometry import mapping as _mapping
    BATCH = 1000
    feature_dicts = []
    for _idx, row in gdf.iterrows():
        key = row.get(section_field)
        if key is None:
            continue
        gj = _json.dumps(_mapping(row.geometry))
        feature_dicts.append({
            "geometry_id": geom_row.id,
            "section_key": str(key),
            "geometry_geojson": gj,
        })
        if len(feature_dicts) >= BATCH:
            db.bulk_insert_mappings(VaisalaNetworkFeature, feature_dicts)
            feature_dicts = []
    if feature_dicts:
        db.bulk_insert_mappings(VaisalaNetworkFeature, feature_dicts)

    db.commit()

    return {
        "id": geom_row.id,
        "source_filename": filename,
        "section_field": section_field,
        "feature_count": geom_row.feature_count,
        "crs_wkt": crs_wkt,
        "fields": fields,
    }


@router.get("/network-geometry/current")
def current_network_geometry(
    authority_id: Optional[int] = Query(None, description="Admin only — filter by authority"),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    auth_id = resolve_authority_id(current_user, authority_id)
    q = db.query(VaisalaNetworkGeometry).filter_by(is_active=True)
    if auth_id is not None:
        q = q.filter(VaisalaNetworkGeometry.authority_id == auth_id)
    geom = q.order_by(VaisalaNetworkGeometry.created_at.desc()).first()
    if not geom:
        return {"active": False}
    return {
        "active": True,
        "id": geom.id,
        "source_filename": geom.source_filename,
        "section_field": geom.section_field,
        "feature_count": geom.feature_count,
        "crs_wkt": geom.crs_wkt,
        "created_at": geom.created_at,
    }


@router.get("/network-geometry/features")
def network_features_for_survey(
    survey_id: int = Query(..., description="Survey to join geometry against"),
    merge_scale: str = Query("section", description="section | 100m | 10m"),
    split: str = Query("combined", description="combined | urban | rural"),
    treatment_mode: str = Query("defect", description="defect | percentile"),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """GeoJSON FeatureCollection: geometry from active network layer joined with survey scores.

    Falls back to zero features if no network layer is active. section_key matches are attempted
    exact first, then with leading-zero normalisation (matches priority_dst.html normaliseSectionId).
    """
    survey = _get_survey(db, survey_id, current_user)
    if not survey:
        raise HTTPException(404, "Survey not found")
    _validate_view_params(merge_scale, split, treatment_mode)

    geom = (
        db.query(VaisalaNetworkGeometry)
          .filter(VaisalaNetworkGeometry.authority_id == survey.authority_id,
                  VaisalaNetworkGeometry.is_active == True)
          .order_by(VaisalaNetworkGeometry.created_at.desc())
          .first()
    )
    if not geom:
        return {"type": "FeatureCollection", "features": [], "network_active": False}

    effective_scale = "section" if split == "urban" else merge_scale
    rows = _build_view_rows(db, survey_id, effective_scale, split)
    if treatment_mode == "percentile":
        _apply_percentile_treatments(rows)

    def _norm(k: str) -> str:
        s = str(k)
        if s.isdigit():
            return s.lstrip("0") or "0"
        return s

    row_by_key: dict = {}
    row_by_norm: dict = {}
    for r in rows:
        ref = str(r.get("section_ref") or "")
        if not ref:
            continue
        row_by_key.setdefault(ref, r)
        row_by_norm.setdefault(_norm(ref), r)

    features = (
        db.query(VaisalaNetworkFeature)
          .filter_by(geometry_id=geom.id)
          .all()
    )

    import json as _json
    out_features: list[dict] = []
    matched = 0
    for f in features:
        key = f.section_key or ""
        r = row_by_key.get(key) or row_by_norm.get(_norm(key))
        if r is not None:
            matched += 1
        try:
            geometry_geojson = _json.loads(f.geometry_geojson)
        except Exception:
            continue
        props = {
            "section_key": key,
            "matched": r is not None,
        }
        if r is not None:
            for k in ("section_ref", "chunk_label", "road_name", "net_reference", "road_class",
                      "urban_rural", "length_m", "priority_score", "worst_interval_score",
                      "rag_band", "treatment",
                      "primary_defect", "primary_defect_contribution",
                      "secondary_defect", "secondary_defect_contribution",
                      "structural_pct", "alligator_pct", "localised_pct",
                      "dressing_pct", "micro_pct", "edge_pct",
                      "road_surface_condition", "asphalt_condition", "pas2161_category"):
                props[k] = r.get(k)
        out_features.append({
            "type": "Feature",
            "geometry": geometry_geojson,
            "properties": props,
        })

    return {
        "type": "FeatureCollection",
        "features": out_features,
        "network_active": True,
        "source_filename": geom.source_filename,
        "section_field": geom.section_field,
        "total_features": geom.feature_count,
        "matched_features": matched,
    }


@router.delete("/network-geometry/{geometry_id}", status_code=204)
def delete_network_geometry(
    geometry_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    geom = db.query(VaisalaNetworkGeometry).filter_by(id=geometry_id, authority_id=current_user.authority_id).first()
    if not geom:
        raise HTTPException(404, "Not found")
    db.delete(geom)
    db.commit()
    return None


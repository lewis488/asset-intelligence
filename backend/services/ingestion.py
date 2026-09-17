"""
Flexible ingestion pipeline for SCANNER Excel, CVI CSV, and reactive jobs CSV.

SCANNER Excel format (WSCC / HMDIF-derived):
  Rows 0–2 are multi-row headers (section group / description / column name).
  Row 3 is a network summary row that must be skipped.
  Rows 4+ are per-section asset data.
  Merged cells in row 0 are forward-filled to give each column its section group.
  B+C sheet has an additional EDI AVERAGE column absent from A ROADS sheet.
"""
import io
import logging
from datetime import datetime
from typing import Optional

import pandas as pd

logger = logging.getLogger(__name__)

# ── SCANNER column mapping ────────────────────────────────────────────────────

# Each entry maps (section_group_keyword, column_keyword) → canonical field name.
# section_group_keyword is matched via substring on the forward-filled group row.

_SECTION_COL_MAP: list[tuple[str, str, str]] = [
    # (group fragment, col fragment, canonical)
    # SECTION INFORMATION
    ("SECTION", "NSG", "nsg_ref"),
    ("SECTION", "ROAD_NAME", "road_name"),
    ("SECTION", "ROAD NAME", "road_name"),
    ("SECTION", "PARISH", "parish"),
    ("SECTION", "LENGTH", "length_m"),
    ("SECTION", "ROAD CLASS", "road_class"),
    ("SECTION", "ROAD_CLASS", "road_class"),
    # DEFECT LENGTHS (% of section)
    ("DEFECT", "OVERALL", "defect_overall_pct"),
    ("DEFECT", "RUTTING", "defect_rutting_pct"),
    ("DEFECT", "CRACKING", "defect_cracking_pct"),
    ("DEFECT", "TEXTURE", "defect_texture_pct"),
    ("DEFECT", "LPV", "defect_lpv_pct"),
    # AMBER LENGTHS ONLY
    ("AMBER", "OVERALL", "amber_length_m"),
    ("AMBER", "RUTTING", "amber_rutting"),
    ("AMBER", "CRACKING", "amber_cracking"),
    ("AMBER", "TEXTURE", "amber_texture"),
    ("AMBER", "LPV", "amber_lpv"),
    # RED LENGTHS ONLY
    ("RED", "OVERALL", "red_length_m"),
    ("RED", "RUTTING", "red_rutting"),
    ("RED", "CRACKING", "red_cracking"),
    ("RED", "TEXTURE", "red_texture"),
    ("RED", "LPV", "red_lpv"),
    # CI DISTRIBUTION (proportional contributions)
    ("CI", "AVERAGE_CI", "avg_ci"),
    ("CI", "AVERAGE CI", "avg_ci"),
    ("CI", "RUTTING", "ci_contribution_rutting"),
    ("CI", "CRACKING", "ci_contribution_cracking"),
    ("CI", "TEXTURE", "ci_contribution_texture"),
    ("CI", "LPV", "ci_contribution_lpv"),
]

# General alias fallback (case-insensitive, applied when section group is unknown)
_GENERAL_ALIASES: dict[str, list[str]] = {
    "nsg_ref": ["NSG", "NSG_REF", "USRN", "ASSET_ID", "ASSET_REF", "SECTION_REF", "SECTION REF"],
    "road_name": ["ROAD_NAME", "ROAD NAME", "NAME", "STREET_NAME", "STREET NAME", "ROAD"],
    "parish": ["PARISH", "AREA", "LOCALITY"],
    "length_m": ["LENGTH", "LENGTH_M", "LEN", "SECTION_LENGTH", "SECTION LENGTH"],
    "avg_ci": ["AVERAGE_CI", "AVERAGE CI", "AVG_CI", "CI", "CONDITION_INDEX", "CI_VALUE"],
    "amber_pct": ["AMBER_PCT", "AMBER%", "%AMBER", "AMBER_PERCENTAGE", "PCT_AMBER"],
    "red_pct": ["RED_PCT", "RED%", "%RED", "RED_PERCENTAGE", "PCT_RED"],
    "edi_avg": ["EDI AVERAGE", "EDI_AVERAGE", "EDI_AVG", "EDI", "EDGE_DETERIORATION_INDEX"],
}


def _detect_road_class(sheet_name: str) -> str:
    su = sheet_name.upper()
    if "A ROAD" in su or su.strip() == "A":
        return "A"
    if "B" in su and "C" in su:
        return "BC"
    return "U"


def _ffill(lst: list) -> list:
    """Forward-fill None values in a list."""
    current = None
    result = []
    for v in lst:
        if v is not None and str(v).strip():
            current = str(v).strip().upper()
        result.append(current)
    return result


def _build_scanner_col_map(groups: list[str], col_names: list) -> tuple[dict, list[str], list[str]]:
    """
    Return (col_index→canonical_name, mapped_descriptions, unmapped_col_names).
    """
    mapping: dict[int, str] = {}
    mapped: list[str] = []
    unmapped: list[str] = []

    for i, (group, col) in enumerate(zip(groups, col_names)):
        # Skip None, empty, or NaN cells (separator columns in the Excel)
        if col is None:
            continue
        try:
            if pd.isna(col):
                continue
        except (TypeError, ValueError):
            pass
        col_s = str(col).strip()
        if not col_s or col_s.lower() == "nan":
            continue

        col_u = col_s.upper()
        group_u = str(group or "").strip().upper()

        # Special case: EDI regardless of group
        if "EDI" in col_u:
            mapping[i] = "edi_avg"
            mapped.append(f"col[{i}] '{col_s}' → edi_avg")
            continue

        # Try section-aware mapping first
        matched = False
        for grp_kw, col_kw, canonical in _SECTION_COL_MAP:
            if grp_kw in group_u and col_kw in col_u:
                if i not in mapping:
                    mapping[i] = canonical
                    mapped.append(f"col[{i}] '{col_s}' ({group}) → {canonical}")
                matched = True
                break

        if not matched:
            # General alias fallback
            for canonical, aliases in _GENERAL_ALIASES.items():
                if col_u in aliases:
                    mapping[i] = canonical
                    mapped.append(f"col[{i}] '{col_s}' → {canonical} (alias)")
                    matched = True
                    break

        if not matched:
            unmapped.append(col_s)  # always a clean string now

    return mapping, mapped, unmapped


def _is_summary_row(row: list, nsg_col_idx: Optional[int]) -> bool:
    """Detect the network summary row (first data row with no real NSG ref)."""
    if nsg_col_idx is None:
        return False
    val = row[nsg_col_idx] if nsg_col_idx < len(row) else None
    if val is None:
        return True
    s = str(val).strip()
    # Summary row typically has the road class or authority name as NSG
    return not any(c.isdigit() for c in s)


def parse_scanner_excel(content: bytes, source_file: str) -> dict:
    """
    Parse a WSCC SCANNER Excel file.
    Returns {"records": [...], "sheets": [{sheet_name, road_class, mapped, unmapped, rows}]}
    """
    try:
        xl = pd.ExcelFile(io.BytesIO(content), engine="openpyxl")
    except Exception as exc:
        raise ValueError(f"Cannot open Excel file: {exc}") from exc

    all_records: list[dict] = []
    sheets_info: list[dict] = []

    for sheet_name in xl.sheet_names:
        try:
            df_raw = pd.read_excel(
                xl,
                sheet_name=sheet_name,
                header=None,
                dtype=object,
            )
        except Exception as exc:
            logger.warning("Skipping sheet '%s': %s", sheet_name, exc)
            continue

        if len(df_raw) < 5:
            logger.warning("Sheet '%s' has fewer than 5 rows — skipping", sheet_name)
            continue

        road_class = _detect_road_class(sheet_name)

        # Forward-fill the section group row (row 0 has merged cell values)
        groups = _ffill(list(df_raw.iloc[0]))
        col_names = list(df_raw.iloc[2])

        col_map, mapped_desc, unmapped = _build_scanner_col_map(groups, col_names)

        if not col_map:
            logger.warning("Sheet '%s': no columns mapped — skipping", sheet_name)
            continue

        # Find NSG column index for summary row detection
        nsg_col = next((i for i, v in col_map.items() if v == "nsg_ref"), None)

        # Data rows: skip header rows (0–2) and summary row (3)
        data_rows = df_raw.iloc[4:].values.tolist()

        ingested = 0
        for row in data_rows:
            if not any(v is not None and str(v).strip() != "" for v in row):
                continue  # skip empty rows
            if _is_summary_row(row, nsg_col):
                continue

            record: dict = {"road_class": road_class, "source_file": source_file}
            for col_idx, field in col_map.items():
                v = row[col_idx] if col_idx < len(row) else None
                # pandas NaT → None
                if pd.isna(v) if not isinstance(v, (list, dict)) else False:
                    v = None
                record[field] = v

            if not record.get("nsg_ref"):
                continue

            record["nsg_ref"] = str(record["nsg_ref"]).strip()
            ingested += 1
            all_records.append(record)

        sheets_info.append({
            "sheet_name": sheet_name,
            "road_class": road_class,
            "mapped_columns": mapped_desc,
            "unmapped_columns": unmapped,
            "rows_ingested": ingested,
        })
        logger.info("Sheet '%s' (%s): %d rows, mapped=%d, unmapped=%d",
                    sheet_name, road_class, ingested, len(mapped_desc), len(unmapped))

    if not all_records:
        raise ValueError(
            "No records parsed from Excel file. "
            "Check sheet names contain 'A ROAD' or 'B' and 'C', "
            "and that NSG/USRN column is present."
        )

    return {"records": all_records, "sheets": sheets_info}


# ── SCANNER flat CSV ──────────────────────────────────────────────────────────

_SCANNER_CSV_ALIASES: dict[str, list[str]] = {
    "nsg_ref":                  ["NSG", "NSG_REF", "USRN", "nsg", "usrn", "asset_ref"],
    "road_name":                ["ROAD_NAME", "road_name", "STREET_NAME", "road"],
    "parish":                   ["PARISH", "parish", "AREA", "area", "locality"],
    "road_class":               ["ROAD_CLASS", "road_class", "CLASS", "class"],
    "length_m":                 ["LENGTH_M", "LENGTH", "length_m", "length", "LEN"],
    "survey_year":              ["SURVEY_YEAR", "survey_year", "YEAR", "year"],
    "avg_ci":                   ["AVG_CI", "AVERAGE_CI", "avg_ci", "CI", "condition_index"],
    "defect_overall_pct":       ["DEFECT_OVERALL_PCT", "defect_overall_pct"],
    "defect_rutting_pct":       ["DEFECT_RUTTING_PCT", "defect_rutting_pct"],
    "defect_cracking_pct":      ["DEFECT_CRACKING_PCT", "defect_cracking_pct"],
    "defect_texture_pct":       ["DEFECT_TEXTURE_PCT", "defect_texture_pct"],
    "defect_lpv_pct":           ["DEFECT_LPV_PCT", "defect_lpv_pct"],
    "amber_length_m":           ["AMBER_LENGTH_M", "amber_length_m"],
    "amber_pct":                ["AMBER_PCT", "amber_pct"],
    "amber_rutting":            ["AMBER_RUTTING", "amber_rutting"],
    "amber_cracking":           ["AMBER_CRACKING", "amber_cracking"],
    "amber_texture":            ["AMBER_TEXTURE", "amber_texture"],
    "amber_lpv":                ["AMBER_LPV", "amber_lpv"],
    "red_length_m":             ["RED_LENGTH_M", "red_length_m"],
    "red_pct":                  ["RED_PCT", "red_pct"],
    "red_rutting":              ["RED_RUTTING", "red_rutting"],
    "red_cracking":             ["RED_CRACKING", "red_cracking"],
    "red_texture":              ["RED_TEXTURE", "red_texture"],
    "red_lpv":                  ["RED_LPV", "red_lpv"],
    "ci_contribution_rutting":  ["CI_CONTRIBUTION_RUTTING",  "ci_contribution_rutting"],
    "ci_contribution_cracking": ["CI_CONTRIBUTION_CRACKING", "ci_contribution_cracking"],
    "ci_contribution_texture":  ["CI_CONTRIBUTION_TEXTURE",  "ci_contribution_texture"],
    "ci_contribution_lpv":      ["CI_CONTRIBUTION_LPV",      "ci_contribution_lpv"],
    "edi_avg":                  ["EDI_AVG", "edi_avg", "EDI"],
}

_SCANNER_CSV_NUMERIC = {
    "length_m", "avg_ci",
    "defect_overall_pct", "defect_rutting_pct", "defect_cracking_pct",
    "defect_texture_pct", "defect_lpv_pct",
    "amber_length_m", "amber_pct", "amber_rutting", "amber_cracking",
    "amber_texture", "amber_lpv",
    "red_length_m", "red_pct", "red_rutting", "red_cracking",
    "red_texture", "red_lpv",
    "ci_contribution_rutting", "ci_contribution_cracking",
    "ci_contribution_texture", "ci_contribution_lpv",
    "edi_avg",
}

# Canonical column order used for the downloadable template
SCANNER_CSV_COLUMNS = list(_SCANNER_CSV_ALIASES.keys())


def parse_scanner_csv(content: bytes, source_file: str) -> dict:
    """
    Parse a flat SCANNER CSV (single header row, one section per data row).
    Returns the same dict structure as parse_scanner_excel so the upload
    router can handle both without branching.

    Column names are matched case-insensitively against the alias table.
    All percentage/contribution values expected as decimals (0.0–1.0).
    CI values as raw numbers (0–150+). Lengths in metres.
    EDI_AVG can be blank for A roads.
    """
    try:
        df = pd.read_csv(io.BytesIO(content))
    except Exception as exc:
        raise ValueError(f"Cannot read SCANNER CSV: {exc}") from exc

    col_lookup = {c.lower().strip(): c for c in df.columns}
    rename: dict[str, str] = {}
    mapped: list[str] = []

    for canonical, aliases in _SCANNER_CSV_ALIASES.items():
        for alias in aliases:
            if alias.lower().strip() in col_lookup:
                orig = col_lookup[alias.lower().strip()]
                if orig not in rename:
                    rename[orig] = canonical
                    mapped.append(f"{orig} → {canonical}")
                break

    unmapped = [c for c in df.columns if c not in rename]
    df = df.rename(columns=rename)

    if "nsg_ref" not in df.columns:
        raise ValueError(
            "Required column 'NSG' / 'NSG_REF' / 'USRN' not found. "
            "Download the CSV template to see the expected column names."
        )

    for col in _SCANNER_CSV_NUMERIC:
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")

    df["source_file"] = source_file
    df = df.dropna(subset=["nsg_ref"])
    df["nsg_ref"] = df["nsg_ref"].astype(str).str.strip()

    records = df.to_dict(orient="records")
    logger.info(
        "SCANNER CSV '%s': %d rows | mapped=%d | unmapped=%d",
        source_file, len(records), len(mapped), len(unmapped),
    )

    return {
        "records": records,
        "sheets": [
            {
                "sheet_name": "CSV Upload",
                "road_class": "mixed",
                "mapped_columns": mapped,
                "unmapped_columns": unmapped,
                "rows_ingested": len(records),
            }
        ],
    }


# ── CVI CSV ───────────────────────────────────────────────────────────────────

_CVI_ALIASES: dict[str, list[str]] = {
    "nsg_ref": ["NSG", "USRN", "nsg_ref", "nsg", "usrn", "asset_id", "ASSET_ID", "section_ref"],
    "road_name": ["ROAD_NAME", "road_name", "name", "NAME", "street_name", "STREET_NAME"],
    "parish": ["PARISH", "parish", "area", "AREA", "locality"],
    "length_m": ["LENGTH", "length_m", "length", "LEN"],
    "structural_ci": ["STRUCTURAL_CI", "structural_ci", "STRUCT_CI", "structural", "CVI_STRUCTURAL"],
    "edge_ci": ["EDGE_CI", "edge_ci", "EDGE", "CVI_EDGE", "edge"],
    "wearing_course_ci": ["WEARING_COURSE_CI", "wearing_course_ci", "WC_CI", "WEARING_COURSE", "CVI_WC", "wearing_course"],
    "survey_year": ["SURVEY_YEAR", "survey_year", "YEAR", "year", "survey_yr"],
}

# BV224b thresholds
_CVI_THRESHOLDS = {"structural_ci": 85, "edge_ci": 50, "wearing_course_ci": 60}


def parse_cvi_csv(content: bytes, source_file: str) -> dict:
    try:
        df = pd.read_csv(io.BytesIO(content))
    except Exception as exc:
        raise ValueError(f"Cannot read CVI CSV: {exc}") from exc

    col_lookup = {c.lower().strip(): c for c in df.columns}
    rename: dict[str, str] = {}
    mapped: list[str] = []

    for canonical, aliases in _CVI_ALIASES.items():
        for alias in aliases:
            if alias.lower().strip() in col_lookup:
                orig = col_lookup[alias.lower().strip()]
                if orig not in rename:
                    rename[orig] = canonical
                    mapped.append(f"{orig} → {canonical}")
                break

    unmapped = [c for c in df.columns if c not in rename]
    df = df.rename(columns=rename)

    required = [f for f in ("nsg_ref",) if f not in df.columns]
    if required:
        raise ValueError(f"Required columns missing: {required}. Check column names.")

    # Apply BV224b flags
    for field, threshold in _CVI_THRESHOLDS.items():
        flag_field = field.replace("_ci", "_flagged")
        if field in df.columns:
            df[field] = pd.to_numeric(df[field], errors="coerce")
            df[flag_field] = df[field].apply(lambda x: bool(x >= threshold) if pd.notna(x) else False)

    df["source_file"] = source_file
    df = df.dropna(subset=["nsg_ref"])
    df["nsg_ref"] = df["nsg_ref"].astype(str).str.strip()

    records = df.to_dict(orient="records")
    logger.info("CVI CSV: %d rows | mapped=%s | unmapped=%s", len(records), mapped, unmapped)
    return {"records": records, "mapped_columns": mapped, "unmapped_columns": unmapped, "row_count": len(records)}


# ── Reactive jobs CSV ─────────────────────────────────────────────────────────

_REACTIVE_ALIASES: dict[str, list[str]] = {
    "nsg_ref": ["NSG", "USRN", "nsg_ref", "nsg", "usrn", "asset_id", "section_ref"],
    "road_name": ["ROAD_NAME", "road_name", "street_name", "STREET_NAME"],
    "job_ref": ["JOB_REF", "job_ref", "ORDER_REF", "order_ref", "job_number", "JOB_NUMBER"],
    "job_type": ["JOB_TYPE", "job_type", "work_type", "WORK_TYPE", "type", "TYPE"],
    "defect_type": ["DEFECT_TYPE", "defect_type", "defect", "DEFECT", "fault_type", "FAULT_TYPE"],
    "job_date": ["JOB_DATE", "job_date", "date", "DATE", "completion_date", "raised_date", "order_date"],
    "cost_gbp": ["COST_GBP", "cost_gbp", "COST", "cost", "amount", "AMOUNT", "actual_cost", "ACTUAL_COST"],
    "response_category": ["RESPONSE_CATEGORY", "response_category", "response", "RESPONSE", "priority", "PRIORITY"],
}


def parse_reactive_csv(content: bytes, source_file: str) -> dict:
    try:
        df = pd.read_csv(io.BytesIO(content))
    except Exception as exc:
        raise ValueError(f"Cannot read reactive CSV: {exc}") from exc

    col_lookup = {c.lower().strip(): c for c in df.columns}
    rename: dict[str, str] = {}
    mapped: list[str] = []

    for canonical, aliases in _REACTIVE_ALIASES.items():
        for alias in aliases:
            if alias.lower().strip() in col_lookup:
                orig = col_lookup[alias.lower().strip()]
                if orig not in rename:
                    rename[orig] = canonical
                    mapped.append(f"{orig} → {canonical}")
                break

    unmapped = [c for c in df.columns if c not in rename]
    df = df.rename(columns=rename)

    if "nsg_ref" not in df.columns:
        raise ValueError("Required column 'NSG/USRN' not found. Check column names.")

    if "job_date" in df.columns:
        df["job_date"] = pd.to_datetime(df["job_date"], errors="coerce", dayfirst=True)
    if "cost_gbp" in df.columns:
        df["cost_gbp"] = pd.to_numeric(df["cost_gbp"], errors="coerce")

    df["source_file"] = source_file
    df = df.dropna(subset=["nsg_ref"])
    df["nsg_ref"] = df["nsg_ref"].astype(str).str.strip()

    records = df.to_dict(orient="records")
    logger.info("Reactive CSV: %d rows | mapped=%s | unmapped=%s", len(records), mapped, unmapped)
    return {"records": records, "mapped_columns": mapped, "unmapped_columns": unmapped, "row_count": len(records)}


def get_alias_reference() -> dict:
    return {
        "scanner_csv": {
            "description": "Flat CSV — one header row, one section per data row. Download template for exact column names.",
            "columns": _SCANNER_CSV_ALIASES,
        },
        "scanner_excel": {
            "description": "Multi-header Excel. Sheet names should contain 'A ROAD' or 'B+C ROAD'.",
            "section_groups": ["SECTION INFORMATION", "DEFECT LENGTHS", "AMBER LENGTHS ONLY", "RED LENGTHS ONLY", "CI DISTRIBUTION"],
            "key_columns": _GENERAL_ALIASES,
        },
        "cvi_csv": _CVI_ALIASES,
        "reactive_csv": _REACTIVE_ALIASES,
    }


# ── Raw Confirm format parsers ────────────────────────────────────────────────
# These parse the direct Confirm export format (10m interval rows for SCANNER/SCRIM,
# variable-length sections for CVI) and aggregate to one record per NSG per survey year.


def _read_tabular(content: bytes) -> pd.DataFrame:
    """Read bytes as DataFrame: try Excel (.xlsx) first, then CSV."""
    try:
        return pd.read_excel(io.BytesIO(content), header=0, dtype=object, engine="openpyxl")
    except Exception:
        try:
            return pd.read_csv(io.BytesIO(content), dtype=object)
        except Exception as exc:
            raise ValueError(f"Cannot read file as Excel or CSV: {exc}") from exc


def _mode_or_none(series: pd.Series):
    """Return the most common non-null value in a Series, or None."""
    s = series.dropna()
    if s.empty:
        return None
    m = s.mode()
    return m.iloc[0] if not m.empty else None


_LARGE_FILE_BYTES = 10 * 1024 * 1024  # 10 MB — threshold for chunked processing
_CHUNK_ROWS = 10_000


def _is_xlsx(content: bytes) -> bool:
    return content[:4] == b"PK\x03\x04"


def _iter_chunks(content: bytes, chunksize: int):
    """Yield DataFrame chunks from Excel or CSV bytes without loading the full file at once."""
    if _is_xlsx(content):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        headers = [str(h) if h is not None else "" for h in next(rows_iter)]
        batch: list = []
        for row in rows_iter:
            batch.append(row)
            if len(batch) >= chunksize:
                yield pd.DataFrame(batch, columns=headers, dtype=object)
                batch = []
        if batch:
            yield pd.DataFrame(batch, columns=headers, dtype=object)
        wb.close()
    else:
        for chunk in pd.read_csv(io.BytesIO(content), dtype=object, chunksize=chunksize):
            yield chunk


def _parse_scanner_raw_chunked(content: bytes, source_file: str) -> list[dict]:
    """
    Memory-efficient chunked aggregation for large SCANNER files (>10 MB).
    Processes _CHUNK_ROWS rows at a time, maintaining running per-group stats
    so the peak DataFrame size is bounded regardless of total file size.
    """
    _known_cols = {
        "NSG", "FEATURE_ID", "ROAD_NAME", "START_CHAINAGE", "END_CHAINAGE",
        "OFFSET", "FEATURE_GROUP", "LANE", "OBSERVATION", "CI_VALUE",
        "SURVEY_NUMBER", "SURVEY_DATE",
    }
    # accumulated[key] holds running stats for each (nsg, year, direction) group
    accumulated: dict[tuple, dict] = {}
    unmapped: list[str] = []

    for chunk_idx, chunk in enumerate(_iter_chunks(content, _CHUNK_ROWS)):
        chunk.columns = [str(c).strip().upper() for c in chunk.columns]

        if chunk_idx == 0:
            required = {"NSG", "CI_VALUE", "SURVEY_DATE", "OFFSET"}
            missing = required - set(chunk.columns)
            if missing:
                raise ValueError(f"SCANNER raw: missing required columns {sorted(missing)}")
            unmapped = [c for c in chunk.columns if c not in _known_cols and not c.startswith("_")]

        chunk["OFFSET"] = pd.to_numeric(chunk["OFFSET"], errors="coerce")
        chunk["CI_VALUE"] = pd.to_numeric(chunk["CI_VALUE"], errors="coerce")
        chunk["NSG"] = chunk["NSG"].astype(str).str.strip()
        chunk["_date"] = pd.to_datetime(chunk["SURVEY_DATE"], dayfirst=True, errors="coerce")
        chunk["_year"] = chunk["_date"].dt.year
        chunk["_dir"] = chunk["OFFSET"].apply(
            lambda x: "nearside" if x == -5 else ("offside" if x == 5 else "centre")
        )

        for (nsg, year, dir_), grp in chunk.groupby(["NSG", "_year", "_dir"], dropna=False):
            if pd.isna(year) or not nsg or nsg == "nan":
                continue

            key = (nsg, int(year), str(dir_))
            ci = grp["CI_VALUE"].dropna()
            n = len(grp)

            # Build partial stats for this chunk-group
            partial: dict = {
                "n": n,
                "ci_sum": float(ci.sum()) if len(ci) > 0 else 0.0,
                "ci_valid_n": len(ci),
                "red_n": int((ci >= 100).sum()),
                "amber_n": int(((ci >= 40) & (ci < 100)).sum()),
                "green_n": int((ci < 40).sum()),
                "ci_min": float(ci.min()) if len(ci) > 0 else None,
                "ci_max": float(ci.max()) if len(ci) > 0 else None,
                "road_name_ctr": {},
                "survey_num_ctr": {},
                "min_date": None,
            }

            if "ROAD_NAME" in grp.columns:
                for v in grp["ROAD_NAME"].dropna():
                    s = str(v)
                    partial["road_name_ctr"][s] = partial["road_name_ctr"].get(s, 0) + 1
            if "SURVEY_NUMBER" in grp.columns:
                for v in grp["SURVEY_NUMBER"].dropna():
                    s = str(v)
                    partial["survey_num_ctr"][s] = partial["survey_num_ctr"].get(s, 0) + 1

            raw_date = grp["_date"].min()
            if not pd.isna(raw_date):
                partial["min_date"] = raw_date

            if key not in accumulated:
                accumulated[key] = partial
            else:
                acc = accumulated[key]
                acc["n"] += partial["n"]
                acc["ci_sum"] += partial["ci_sum"]
                acc["ci_valid_n"] += partial["ci_valid_n"]
                acc["red_n"] += partial["red_n"]
                acc["amber_n"] += partial["amber_n"]
                acc["green_n"] += partial["green_n"]
                if partial["ci_min"] is not None:
                    acc["ci_min"] = min(acc["ci_min"], partial["ci_min"]) if acc["ci_min"] is not None else partial["ci_min"]
                if partial["ci_max"] is not None:
                    acc["ci_max"] = max(acc["ci_max"], partial["ci_max"]) if acc["ci_max"] is not None else partial["ci_max"]
                for k, v in partial["road_name_ctr"].items():
                    acc["road_name_ctr"][k] = acc["road_name_ctr"].get(k, 0) + v
                for k, v in partial["survey_num_ctr"].items():
                    acc["survey_num_ctr"][k] = acc["survey_num_ctr"].get(k, 0) + v
                if partial["min_date"] is not None:
                    if acc["min_date"] is None or partial["min_date"] < acc["min_date"]:
                        acc["min_date"] = partial["min_date"]

    if not accumulated:
        raise ValueError(
            "SCANNER raw: no records parsed. "
            "Check that NSG, CI_VALUE, SURVEY_DATE, OFFSET columns are present."
        )

    records: list[dict] = []
    for (nsg, year, dir_), acc in accumulated.items():
        n = acc["n"]
        ci_valid_n = acc["ci_valid_n"]

        if ci_valid_n > 0:
            avg_ci_val = acc["ci_sum"] / ci_valid_n
            red_pct = acc["red_n"] / n * 100
            amber_pct = acc["amber_n"] / n * 100
            green_pct = acc["green_n"] / n * 100
            max_ci_val = acc["ci_max"]
            min_ci_val = acc["ci_min"]
        else:
            avg_ci_val = red_pct = amber_pct = green_pct = max_ci_val = min_ci_val = None

        if red_pct and red_pct > 0:
            rci_band = "Red"
        elif amber_pct and amber_pct > 0:
            rci_band = "Amber"
        else:
            rci_band = "Green"

        rn_ctr = acc["road_name_ctr"]
        road_name = max(rn_ctr, key=rn_ctr.get) if rn_ctr else None
        sn_ctr = acc["survey_num_ctr"]
        survey_number = max(sn_ctr, key=sn_ctr.get) if sn_ctr else None

        raw_date = acc["min_date"]
        survey_date = raw_date.to_pydatetime() if raw_date is not None and hasattr(raw_date, "to_pydatetime") else None

        records.append({
            "nsg_ref": nsg,
            "road_name": road_name,
            "survey_year": year,
            "survey_date": survey_date,
            "survey_number": survey_number,
            "total_length_m": float(n * 10),
            "avg_ci": avg_ci_val,
            "red_pct": red_pct,
            "amber_pct": amber_pct,
            "green_pct": green_pct,
            "max_ci": max_ci_val,
            "min_ci": min_ci_val,
            "rci_band": rci_band,
            "offset_direction": dir_,
            "source_file": source_file,
            "_unmapped_columns": unmapped,
        })

    logger.info("SCANNER raw (chunked) '%s': %d aggregated records", source_file, len(records))
    return records


def parse_scanner_raw(content: bytes, source_file: str) -> list[dict]:
    """
    Parse a Confirm SCANNER export (raw 10m interval rows).

    Aggregates to one record per (NSG, survey_year, offset_direction).
    CI bands: Green < 40, Amber 40-99, Red >= 100.
    Each row represents a 10m section; total_length_m = row_count * 10.
    """
    if len(content) > _LARGE_FILE_BYTES:
        return _parse_scanner_raw_chunked(content, source_file)

    df = _read_tabular(content)
    df.columns = [str(c).strip().upper() for c in df.columns]

    required = {"NSG", "CI_VALUE", "SURVEY_DATE", "OFFSET"}
    missing = required - set(df.columns)
    if missing:
        raise ValueError(f"SCANNER raw: missing required columns {sorted(missing)}")

    df["OFFSET"] = pd.to_numeric(df["OFFSET"], errors="coerce")
    df["CI_VALUE"] = pd.to_numeric(df["CI_VALUE"], errors="coerce")
    df["NSG"] = df["NSG"].astype(str).str.strip()
    df["_survey_date"] = pd.to_datetime(df["SURVEY_DATE"], dayfirst=True, errors="coerce")
    df["_survey_year"] = df["_survey_date"].dt.year
    df["_offset_dir"] = df["OFFSET"].apply(
        lambda x: "nearside" if x == -5 else ("offside" if x == 5 else "centre")
    )

    _known_cols = {
        "NSG", "FEATURE_ID", "ROAD_NAME", "START_CHAINAGE", "END_CHAINAGE",
        "OFFSET", "FEATURE_GROUP", "LANE", "OBSERVATION", "CI_VALUE",
        "SURVEY_NUMBER", "SURVEY_DATE",
    }
    unmapped = [c for c in df.columns if c not in _known_cols and not c.startswith("_")]

    records: list[dict] = []
    for (nsg, year, dir_), grp in df.groupby(["NSG", "_survey_year", "_offset_dir"], dropna=False):
        if pd.isna(year) or not nsg or nsg == "nan":
            continue

        ci = grp["CI_VALUE"].dropna()
        n = len(grp)

        road_name_raw = _mode_or_none(grp["ROAD_NAME"]) if "ROAD_NAME" in grp.columns else None
        road_name = str(road_name_raw) if road_name_raw is not None else None

        survey_number_raw = _mode_or_none(grp["SURVEY_NUMBER"]) if "SURVEY_NUMBER" in grp.columns else None
        survey_number = str(survey_number_raw) if survey_number_raw is not None else None

        raw_date = grp["_survey_date"].min()
        survey_date = raw_date.to_pydatetime() if not pd.isna(raw_date) else None

        if len(ci) > 0:
            avg_ci_val = float(ci.mean())
            red_count = int((ci >= 100).sum())
            amber_count = int(((ci >= 40) & (ci < 100)).sum())
            green_count = int((ci < 40).sum())
            red_pct = red_count / n * 100
            amber_pct = amber_count / n * 100
            green_pct = green_count / n * 100
            max_ci_val = float(ci.max())
            min_ci_val = float(ci.min())
        else:
            avg_ci_val = red_pct = amber_pct = green_pct = max_ci_val = min_ci_val = None

        if red_pct and red_pct > 0:
            rci_band = "Red"
        elif amber_pct and amber_pct > 0:
            rci_band = "Amber"
        else:
            rci_band = "Green"

        records.append({
            "nsg_ref": nsg,
            "road_name": road_name,
            "survey_year": int(year),
            "survey_date": survey_date,
            "survey_number": survey_number,
            "total_length_m": float(n * 10),
            "avg_ci": avg_ci_val,
            "red_pct": red_pct,
            "amber_pct": amber_pct,
            "green_pct": green_pct,
            "max_ci": max_ci_val,
            "min_ci": min_ci_val,
            "rci_band": rci_band,
            "offset_direction": dir_,
            "source_file": source_file,
            "_unmapped_columns": unmapped,
        })

    if not records:
        raise ValueError(
            "SCANNER raw: no records parsed. "
            "Check that NSG, CI_VALUE, SURVEY_DATE, OFFSET columns are present."
        )

    logger.info("SCANNER raw '%s': %d aggregated records", source_file, len(records))
    return records


def parse_cvi_raw(content: bytes, source_file: str) -> list[dict]:
    """
    Parse a Confirm CVI export (variable-length section rows).

    Aggregates to one record per (NSG, survey_year).
    Overall CI is length-weighted. Domain CIs use maximum (BV224b flag if ANY section exceeds threshold).
    BV224b thresholds: structural >= 85, edge >= 50, wearing course >= 60.
    """
    df = _read_tabular(content)
    df.columns = [str(c).strip().upper() for c in df.columns]

    required = {"NSG", "CI_OVRLL", "SECTIONLEN", "SURVEY_DATE"}
    missing = required - set(df.columns)
    if missing:
        raise ValueError(f"CVI raw: missing required columns {sorted(missing)}")

    for col in ("CI_STRUC", "CI_WCRSE", "CI_EDGE", "CI_OVRLL", "SECTIONLEN"):
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")

    df["NSG"] = df["NSG"].astype(str).str.strip()
    df["_survey_date"] = pd.to_datetime(df["SURVEY_DATE"], dayfirst=True, errors="coerce")
    df["_survey_year"] = df["_survey_date"].dt.year

    _known_cols = {
        "NSG", "FEATURE_ID", "ROAD_NAME", "START_CHAINAGE", "END_CHAINAGE",
        "CI_STRUC", "CI_WCRSE", "CI_EDGE", "CI_OVRLL", "PARISH", "SECTIONLEN",
        "SURVEY_NUMBER", "SURVEY_DATE", "SURVEY_NAME", "OFFSET",
    }
    unmapped = [c for c in df.columns if c not in _known_cols and not c.startswith("_")]

    records: list[dict] = []
    for (nsg, year), grp in df.groupby(["NSG", "_survey_year"], dropna=False):
        if pd.isna(year) or not nsg or nsg == "nan":
            continue

        slen = grp["SECTIONLEN"].fillna(0)
        total_length = float(slen.sum())

        road_name_raw = _mode_or_none(grp["ROAD_NAME"]) if "ROAD_NAME" in grp.columns else None
        road_name = str(road_name_raw) if road_name_raw is not None else None

        parish_raw = _mode_or_none(grp["PARISH"]) if "PARISH" in grp.columns else None
        parish = str(parish_raw) if parish_raw is not None else None

        survey_name_raw = _mode_or_none(grp["SURVEY_NAME"]) if "SURVEY_NAME" in grp.columns else None
        survey_name = str(survey_name_raw) if survey_name_raw is not None else None

        raw_date = grp["_survey_date"].min()
        survey_date = raw_date.to_pydatetime() if not pd.isna(raw_date) else None

        # Length-weighted overall CI
        ci_ovrll = grp["CI_OVRLL"].fillna(0)
        avg_ci_overall = float((ci_ovrll * slen).sum() / total_length) if total_length > 0 else None

        # Max per domain (any section exceeding threshold triggers BV224b flag)
        def _max_if_present(col: str):
            if col in grp.columns and grp[col].notna().any():
                return float(grp[col].max())
            return None

        max_ci_struct = _max_if_present("CI_STRUC")
        max_ci_edge = _max_if_present("CI_EDGE")
        max_ci_wc = _max_if_present("CI_WCRSE")

        structural_flagged = bool(max_ci_struct is not None and max_ci_struct >= 85)
        edge_flagged = bool(max_ci_edge is not None and max_ci_edge >= 50)
        wc_flagged = bool(max_ci_wc is not None and max_ci_wc >= 60)
        any_flagged = structural_flagged or edge_flagged or wc_flagged

        # % of total length where each domain exceeds threshold
        def _pct_flagged_length(col: str, threshold: float) -> Optional[float]:
            if col not in grp.columns or total_length <= 0:
                return None
            mask = grp[col].fillna(0) >= threshold
            return float(slen[mask].sum() / total_length * 100)

        records.append({
            "nsg_ref": nsg,
            "road_name": road_name,
            "parish": parish,
            "survey_year": int(year),
            "survey_date": survey_date,
            "survey_name": survey_name,
            "total_length_m": total_length,
            "avg_ci_overall": avg_ci_overall,
            "max_ci_structural": max_ci_struct,
            "max_ci_edge": max_ci_edge,
            "max_ci_wearingcourse": max_ci_wc,
            "structural_flagged": structural_flagged,
            "edge_flagged": edge_flagged,
            "wearingcourse_flagged": wc_flagged,
            "any_flagged": any_flagged,
            "pct_length_structural_flagged": _pct_flagged_length("CI_STRUC", 85),
            "pct_length_edge_flagged": _pct_flagged_length("CI_EDGE", 50),
            "pct_length_wc_flagged": _pct_flagged_length("CI_WCRSE", 60),
            "source_file": source_file,
            "_unmapped_columns": unmapped,
        })

    if not records:
        raise ValueError(
            "CVI raw: no records parsed. "
            "Check that NSG, CI_OVRLL, SECTIONLEN, SURVEY_DATE columns are present."
        )

    logger.info("CVI raw '%s': %d aggregated records", source_file, len(records))
    return records


def parse_scrim_raw(content: bytes, source_file: str) -> list[dict]:
    """
    Parse a Confirm SCRIM export (10m interval rows, MSSC already averaged across seasons).

    Aggregates to one record per (NSG, survey_year).
    SFC=0 rows excluded from mean/min (zero = no reading, not actual zero skid resistance).
    safety_flagged=True if any section has XDIF < 0 (below investigatory level).
    """
    df = _read_tabular(content)
    df.columns = [str(c).strip().upper() for c in df.columns]

    required = {"NSG", "SFC", "SURVEYDATE"}
    missing = required - set(df.columns)
    if missing:
        raise ValueError(f"SCRIM raw: missing required columns {sorted(missing)}")

    for col in ("SFC", "SFCT", "XDIF"):
        if col in df.columns:
            df[col] = pd.to_numeric(df[col], errors="coerce")

    df["NSG"] = df["NSG"].astype(str).str.strip()
    df["_survey_date"] = pd.to_datetime(df["SURVEYDATE"], dayfirst=True, errors="coerce")
    df["_survey_year"] = df["_survey_date"].dt.year

    _known_cols = {
        "NSG", "FEATURE_ID", "ROAD_NAME", "PARISH", "ILCT", "SFC", "SFCT", "XDIF",
        "STARTCHAIN", "ENDCHAIN", "XSP_CODE", "OFFSET", "SURVEYNUMB",
        "SURVEYDATE", "SURVEY_NAME",
    }
    unmapped = [c for c in df.columns if c not in _known_cols and not c.startswith("_")]

    records: list[dict] = []
    for (nsg, year), grp in df.groupby(["NSG", "_survey_year"], dropna=False):
        if pd.isna(year) or not nsg or nsg == "nan":
            continue

        n = len(grp)

        road_name_raw = _mode_or_none(grp["ROAD_NAME"]) if "ROAD_NAME" in grp.columns else None
        road_name = str(road_name_raw) if road_name_raw is not None else None

        parish_raw = _mode_or_none(grp["PARISH"]) if "PARISH" in grp.columns else None
        parish = str(parish_raw) if parish_raw is not None else None

        survey_name_raw = _mode_or_none(grp["SURVEY_NAME"]) if "SURVEY_NAME" in grp.columns else None
        survey_name = str(survey_name_raw) if survey_name_raw is not None else None

        survey_number_raw = _mode_or_none(grp["SURVEYNUMB"]) if "SURVEYNUMB" in grp.columns else None
        survey_number = str(survey_number_raw) if survey_number_raw is not None else None

        dominant_ilct_raw = _mode_or_none(grp["ILCT"]) if "ILCT" in grp.columns else None
        dominant_ilct = str(dominant_ilct_raw) if dominant_ilct_raw is not None else None

        sfct_raw = _mode_or_none(grp["SFCT"]) if "SFCT" in grp.columns else None
        sfct_threshold = float(sfct_raw) if sfct_raw is not None and not pd.isna(sfct_raw) else None

        raw_date = grp["_survey_date"].min()
        survey_date = raw_date.to_pydatetime() if not pd.isna(raw_date) else None

        # SFC stats — exclude zero readings (0 = no data, not actual zero skid)
        if "SFC" in grp.columns:
            sfc_valid = grp["SFC"][grp["SFC"] > 0].dropna()
            mean_sfc = float(sfc_valid.mean()) if not sfc_valid.empty else None
            min_sfc = float(sfc_valid.min()) if not sfc_valid.empty else None
        else:
            mean_sfc = min_sfc = None

        # XDIF stats
        if "XDIF" in grp.columns:
            xdif = grp["XDIF"]
            sections_below_il = int((xdif < 0).sum())
            pct_below_il = float(sections_below_il / n * 100)
            safety_flagged = sections_below_il > 0
            worst_xdif = float(xdif.min()) if xdif.notna().any() else None
        else:
            sections_below_il = 0
            pct_below_il = 0.0
            safety_flagged = False
            worst_xdif = None

        records.append({
            "nsg_ref": nsg,
            "road_name": road_name,
            "parish": parish,
            "survey_year": int(year),
            "survey_date": survey_date,
            "survey_name": survey_name,
            "survey_number": survey_number,
            "total_sections": n,
            "dominant_ilct": dominant_ilct,
            "sfct_threshold": sfct_threshold,
            "mean_sfc": mean_sfc,
            "min_sfc": min_sfc,
            "sections_below_il": sections_below_il,
            "pct_below_il": pct_below_il,
            "safety_flagged": safety_flagged,
            "worst_xdif": worst_xdif,
            "source_file": source_file,
            "_unmapped_columns": unmapped,
        })

    if not records:
        raise ValueError(
            "SCRIM raw: no records parsed. "
            "Check that NSG, SFC, SURVEYDATE columns are present."
        )

    logger.info("SCRIM raw '%s': %d aggregated records", source_file, len(records))
    return records


# ── Raw reactive jobs parser ──────────────────────────────────────────────────

_REACTIVE_INCLUDE_TYPES_LOWER: list[str] = [
    "potholes | cway",
    "patching | cway",
    "verge repairs",
    "kerb | edge works",
    "covers | gullies cway",
]

_REACTIVE_VALID_STATUSES: set[str] = {"works complete", "work inspected"}


def _reactive_is_condition_relevant(job_type_name: str) -> bool:
    lower = str(job_type_name).lower().strip()
    return any(t in lower for t in _REACTIVE_INCLUDE_TYPES_LOWER)


def _reactive_is_valid_status(status_name: str) -> bool:
    return str(status_name).lower().strip() in _REACTIVE_VALID_STATUSES


def _reactive_priority_category(priority_name: str) -> Optional[int]:
    """Extract numeric category from 'Cat1: 2_Hours' → 1, 'Cat3: 5 Days' → 3, etc."""
    s = str(priority_name or "").lower()
    for i in range(1, 6):
        if f"cat{i}" in s:
            return i
    return None


def _reactive_job_type_category(job_type_name: str) -> str:
    jt = str(job_type_name).lower()
    if "pothole" in jt:
        return "pothole"
    if "patching" in jt:
        return "patching"
    if "verge" in jt or "kerb" in jt or "edge" in jt:
        return "edge"
    if "gull" in jt or "drainage" in jt or "cover" in jt:
        return "drainage"
    return "other"


def parse_reactive_raw(content: bytes, source_file: str) -> dict:
    """
    Parse a Confirm reactive jobs export, filtering to condition-relevant types
    and valid completion statuses.

    Condition-relevant types (substring match, case-insensitive):
      Potholes | CWAY, Patching | CWAY, Verge Repairs, Kerb | Edge Works,
      Covers | Gullies CWAY.
    Valid statuses: Works Complete, Work Inspected.

    Returns {"records": [...], "filtered_out": N, "job_type_breakdown": {...}}
    Each record dict has all fields needed to build a ReactiveJobRecord row.
    """
    df = _read_tabular(content)
    df.columns = [str(c).strip().lower() for c in df.columns]

    # NSG ref — Confirm exports use site_code
    if "site_code" in df.columns:
        df["_nsg"] = df["site_code"].astype(str).str.strip()
    elif "nsg" in df.columns:
        df["_nsg"] = df["nsg"].astype(str).str.strip()
    else:
        raise ValueError("Reactive raw: 'site_code' (or 'nsg') column not found.")

    for req in ("job_type_name", "status_name"):
        if req not in df.columns:
            raise ValueError(f"Reactive raw: required column '{req}' not found.")

    total_rows = len(df)

    keep = (
        df["job_type_name"].apply(_reactive_is_condition_relevant) &
        df["status_name"].apply(_reactive_is_valid_status)
    )
    df_valid = df[keep].copy()
    filtered_out = total_rows - len(df_valid)

    if df_valid.empty:
        logger.warning("Reactive raw '%s': all %d rows filtered out", source_file, total_rows)
        return {"records": [], "filtered_out": filtered_out, "job_type_breakdown": {}}

    # Parse dates
    for col in ("job_entry_date", "actual_comp_date"):
        if col in df_valid.columns:
            df_valid[col] = pd.to_datetime(df_valid[col], dayfirst=True, errors="coerce")

    # Derived columns
    if "priority_name" in df_valid.columns:
        df_valid["_priority_cat"] = df_valid["priority_name"].apply(_reactive_priority_category)
    else:
        df_valid["_priority_cat"] = None

    df_valid["_job_type_cat"] = df_valid["job_type_name"].apply(_reactive_job_type_category)

    for col in ("easting", "northing"):
        if col in df_valid.columns:
            df_valid[col] = pd.to_numeric(df_valid[col], errors="coerce")

    def _s(row, col: str) -> Optional[str]:
        v = row.get(col)
        if v is None:
            return None
        try:
            if pd.isna(v):
                return None
        except (TypeError, ValueError):
            pass
        s = str(v).strip()
        return s if s and s.lower() != "nan" else None

    def _dt(row, col: str):
        v = row.get(col)
        if v is None:
            return None
        try:
            if pd.isna(v):
                return None
        except (TypeError, ValueError):
            pass
        return v.to_pydatetime() if hasattr(v, "to_pydatetime") else v

    def _f(row, col: str) -> Optional[float]:
        v = row.get(col)
        if v is None:
            return None
        try:
            f = float(v)
            return None if pd.isna(f) else f
        except (TypeError, ValueError):
            return None

    records: list[dict] = []
    for _, row in df_valid.iterrows():
        nsg = row["_nsg"]
        if not nsg or nsg == "nan":
            continue

        pc = row.get("_priority_cat")
        try:
            priority_cat = int(pc) if pc is not None and not pd.isna(pc) else None
        except (TypeError, ValueError):
            priority_cat = None

        records.append({
            "nsg_ref": nsg,
            "road_name": _s(row, "site_name"),
            "job_number": _s(row, "job_number"),
            "job_entry_date": _dt(row, "job_entry_date"),
            "actual_comp_date": _dt(row, "actual_comp_date"),
            "priority_name": _s(row, "priority_name"),
            "priority_category": priority_cat,
            "job_type_name": _s(row, "job_type_name"),
            "job_type_category": row["_job_type_cat"],
            "status_name": _s(row, "status_name"),
            "district_name": _s(row, "district_name"),
            "locality_name": _s(row, "locality_name"),
            "town_name": _s(row, "town_name"),
            "road_class": _s(row, "road_class"),
            "easting": _f(row, "easting"),
            "northing": _f(row, "northing"),
            "source_file": source_file,
        })

    breakdown: dict[str, int] = {}
    for rec in records:
        cat = rec.get("job_type_category") or "other"
        breakdown[cat] = breakdown.get(cat, 0) + 1

    logger.info(
        "Reactive raw '%s': %d records kept, %d filtered out | breakdown: %s",
        source_file, len(records), filtered_out, breakdown,
    )
    return {"records": records, "filtered_out": filtered_out, "job_type_breakdown": breakdown}


# ── Network shapefile / geopackage ingestion ──────────────────────────────────

_NETWORK_VALID_CLASSES = {"A", "B", "C", "D"}
_NETWORK_OWNER = "WEST SUSSEX COUNTY COUNCIL"


def parse_network_file(content: bytes, source_file: str) -> list[dict]:
    """
    Read a WSCC road network file (.gpkg or .zip containing .shp).
    Filters to ABCD classes + WSCC ownership, aggregates to one row per NSG,
    reprojects geometry to WGS84, returns list of dicts ready for DB insert.
    """
    import io
    import os
    import tempfile
    import zipfile

    import geopandas as gpd
    from shapely.ops import unary_union

    # Write content to a temp file so geopandas can read it
    suffix = ".zip" if source_file.lower().endswith(".zip") else ".gpkg"
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
        tmp.write(content)
        tmp_path = tmp.name

    try:
        if source_file.lower().endswith(".zip"):
            # Unzip and read the .shp inside
            with tempfile.TemporaryDirectory() as tmpdir:
                with zipfile.ZipFile(tmp_path) as zf:
                    zf.extractall(tmpdir)
                shp_files = [f for f in os.listdir(tmpdir) if f.lower().endswith(".shp")]
                if not shp_files:
                    raise ValueError("No .shp file found in ZIP archive")
                gdf = gpd.read_file(os.path.join(tmpdir, shp_files[0]))
        else:
            # GeoPackage — detect layer name using sqlite3 (no fiona required)
            import sqlite3 as _sqlite3
            conn = _sqlite3.connect(tmp_path)
            _cur = conn.cursor()
            _cur.execute("SELECT table_name FROM gpkg_contents WHERE data_type='features'")
            _layers = [r[0] for r in _cur.fetchall()]
            conn.close()
            layer = "Road_Network" if "Road_Network" in _layers else (_layers[0] if _layers else None)
            gdf = gpd.read_file(tmp_path, layer=layer) if layer else gpd.read_file(tmp_path)
    finally:
        os.unlink(tmp_path)

    # Apply filters
    gdf = gdf[
        gdf["CLASS"].isin(_NETWORK_VALID_CLASSES) &
        (gdf["OWNERSHIP"] == _NETWORK_OWNER)
    ].copy()

    if gdf.empty:
        raise ValueError("No features remain after CLASS and OWNERSHIP filters")

    # Reproject to WGS84
    if gdf.crs is None or gdf.crs.to_epsg() != 4326:
        gdf = gdf.to_crs(epsg=4326)

    def _mode(series):
        vc = series.dropna().value_counts()
        return str(vc.index[0]) if len(vc) > 0 else None

    records = []
    for nsg, grp in gdf.groupby("NSGNO"):
        geom = unary_union(grp.geometry.values)
        records.append({
            "nsg_ref":        str(nsg).lstrip('0') or '0',
            "wsccnet":        _mode(grp["WSCCNET"]),
            "road_name":      _mode(grp["ROADNAME"]),
            "road_class":     _mode(grp["CLASS"]),
            "parish":         _mode(grp["PARISH"]),
            "locality":       _mode(grp["LOCALITY"]),
            "district":       _mode(grp["DISTRICT"]),
            "urban_rural":    _mode(grp["URBANRURAL"]),
            "speed_limit":    _mode(grp["SPDLIMIT"]),
            "inspection_freq": _mode(grp["CWINSPFRQ"]),
            "length_m":       float(grp["CWLENGTH"].sum()),
            "avg_width_m":    float(grp["CWAVGWDTH"].mean()),
            "area_m2":        float(grp["CWAREA"].sum()),
            "geometry":       geom.wkt if geom and not geom.is_empty else None,
            "source_file":    source_file,
        })

    logger.info(
        "Network '%s': %d features → %d unique NSGs after filter+aggregate",
        source_file, len(gdf), len(records),
    )
    return records

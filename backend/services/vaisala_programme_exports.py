"""Spreadsheet handover from the same deterministic payload displayed by the API."""
import csv
import io
import json

from openpyxl import Workbook

from services.vaisala_programme import ACTIONS

METHOD_NOTES = [
    "This is an evidence-led worklist, not an approved construction programme.",
    "Lengths are assessed coverage, not repair quantities. Queue lengths may overlap; do not sum them as unique network coverage.",
    "Condition ranks are separate by survey, effective scale and primary model action. Equal scores share rank; unknown scores are unranked.",
    "Validation order is not condition urgency. Monitoring/no-action queues have no urgency rank.",
    "Condition percentiles are retained from the full survey/scale cohort before filtering; they do not determine treatment suitability.",
    "Client decisions are separate from the frozen model recommendation. Completion does not establish repaired condition.",
    "Action worksheets follow the current client action when one is recorded. The exported queue rank remains the original model-action rank, not a rank in the client-selected queue.",
    "Survey date is unknown unless present in source evidence. Import and generation times are not survey dates.",
    "Section references are supplied source references, not automatically confirmed NSG links. No geometry is required for this handover.",
]

FIELDS = [
    "item_key", "survey_id", "section_ref", "net_reference", "road_name", "parent_section_id",
    "narrative_section_id", "source_interval_ids", "assessment_scope", "urban_rural", "from_m", "to_m",
    "assessed_length_m", "length_basis", "recommended_action", "effective_action", "client_action_overridden", "action_label", "brief", "next_question",
    "reason_codes", "prerequisite_tasks", "evidence_status", "priority_score", "rag_band", "queue_rank",
    "queue_size", "validation_order", "priority_explanation", "priority_percentile", "priority_cohort_size",
    "qc_completeness_pct", "qc_reliability_pct", "defect_proportions", "treatment_candidates",
    "treatment_assessment", "evidence_flags", "limitations", "review_status", "client_action", "comment",
    "assignee", "review_sequence", "reviewer_id", "reviewer_name", "reviewed_at", "model_version",
    "policy_version", "programme_id", "generated_at", "export_scope", "export_filters",
    "source_evidence_json", "review_history_json",
]


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, (dict, list, tuple)):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True, allow_nan=False)
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    if isinstance(value, str) and value.startswith(("\t", "\r", "\n")):
        return "'" + value
    return value


def _flat_item(item, programme):
    review = item.get("review") or {}
    result = {**item,
        "effective_action": review.get("client_action") or item["recommended_action"],
        "survey_id": programme["survey_id"],
        "action_label": ACTIONS.get(item["recommended_action"], item["recommended_action"]),
        "review_status": review.get("status", "unreviewed"),
        "client_action": review.get("client_action"), "comment": review.get("comment"),
        "assignee": review.get("assignee"), "review_sequence": review.get("sequence", 0),
        "reviewer_id": review.get("reviewer_id"), "reviewer_name": review.get("reviewer_name"),
        "reviewed_at": review.get("created_at"),
        "model_version": programme["model_version"], "policy_version": programme["policy_version"],
        "programme_id": programme.get("id"), "generated_at": programme.get("generated_at"),
        "export_scope": "filtered" if programme.get("export_filters") else "whole programme",
        "export_filters": programme.get("export_filters", {}),
        "source_evidence_json": {k: v for k, v in item.items() if k not in ("review", "review_history")},
        "review_history_json": item.get("review_history", []),
    }
    return [_cell(result.get(field)) for field in FIELDS]


def export_programme(programme: dict, *, format: str) -> bytes:
    if format == "csv":
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow(FIELDS)
        for item in programme["items"]:
            writer.writerow(_flat_item(item, programme))
        return output.getvalue().encode("utf-8-sig")
    if format != "xlsx":
        raise ValueError("format must be csv or xlsx")

    workbook = Workbook(write_only=True)
    summary = workbook.create_sheet("Summary")
    summary.append(["Programme field", "Value"])
    for key in ("id", "survey_id", "model_version", "policy_version", "generated_at", "merge_scale", "split", "source_filename", "imported_at"):
        summary.append([key, _cell(programme.get(key))])
    for key, value in programme["summary"].items():
        summary.append([key, _cell(value)])
    summary.append(["Export scope", "filtered" if programme.get("export_filters") else "whole programme"])
    summary.append(["Exported records", len(programme["items"])])
    summary.append(["Filters", _cell(programme.get("export_filters", {}))])
    summary.append(["Totals scope", "Whole programme totals; filters only change exported item rows."])

    all_items = workbook.create_sheet("All items")
    all_items.append(FIELDS)
    sheet_names = {"engineer_assessment": "Engineer assessment", "evidence_validation": "Evidence validation",
                   "treatment_appraisal": "Treatment appraisal", "monitor": "Monitor", "no_action_indicated": "No action indicated"}
    action_sheets = {}
    for action, title in sheet_names.items():
        sheet = workbook.create_sheet(title)
        sheet.append(FIELDS)
        action_sheets[action] = sheet
    overflow = []
    for item in programme["items"]:
        row = _flat_item(item, programme)
        for index, value in enumerate(row):
            if isinstance(value, str) and len(value) > 32767:
                overflow.extend([item['item_key'], FIELDS[index], part // 30000 + 1, value[part:part + 30000]]
                                for part in range(0, len(value), 30000))
                row[index] = 'See Evidence continuation; concatenate parts in order for this item and field'
        all_items.append(row)
        action_sheets[item.get("review", {}).get("client_action") or item["recommended_action"]].append(row)
    methodology = workbook.create_sheet("Methodology")
    for note in METHOD_NOTES:
        methodology.append([note])
    methodology.append(["Policy", _cell(programme.get("policy", {}))])
    methodology.append(["Cohorts", _cell(programme.get("cohorts", []))])
    locations = workbook.create_sheet('Local defect review')
    locations.append(['Item key', 'NSG section', 'Defect', 'Sub-section reference', 'From (m)', 'To (m)',
                      'Interval measure (%)', 'Location status'])
    for item in programme['items']:
        for flag in (item.get('treatment_assessment') or {}).get('local_defect_flags') or []:
            for location in flag.get('locations') or [{}]:
                locations.append([_cell(value) for value in [item['item_key'], item.get('section_ref'),
                    flag['defect'], location.get('net_reference'), location.get('from_m'), location.get('to_m'),
                    location.get('interval_measure_pct'), flag.get('location_status')]])
    if overflow:
        from openpyxl.cell import WriteOnlyCell
        continuation = workbook.create_sheet('Evidence continuation')
        continuation.append(['Item key', 'Field', 'Part', 'Text (concatenate parts)'])
        for row in overflow:
            text_cell = WriteOnlyCell(continuation, value=row[3])
            text_cell.data_type = 's'
            continuation.append([_cell(value) for value in row[:3]] + [text_cell])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
